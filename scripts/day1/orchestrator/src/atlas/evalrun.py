"""The phased evaluation runner: the seat test sets (D-49, T-27) and the DeepSeek V4.1 trial (D-46, T-30).

Nothing on Day 1 runs it. The sets do not exist yet: Phase 3 step 2 raises the to-dos `trial-task-set` and
`seat-test-sets` (phase3/loadtest.py:324-325) and only the Principal closes them. The design is the one the Principal
approved on 2026-10-17 (research_notes/ATLAS open inputs 2026-10-17/recommendations.json, items D-49 and D-46,
"build_changes_if_approved"; notes_SETS.md sections 3 to 6; notes_verify_COMP_SETS.md). Run it as root, as
`atlas-admin eval ...` (admin.py) or `python -m atlas.evalrun ...`:

  validate SET.jsonl   the strict schema check, then an advisory report against the approved counts
  freeze SET.jsonl     copy the set into a new run, write the blinding key (0600) and record in the ledger, BEFORE
                       anything runs, the set's sha256, the key's sha256 and the pre-registered decision rule
  run RUN_ID           answers, then judges (--phase), one engine at a time, through the Engine Arbiter
  workbook RUN_ID      the blinded scoring workbook for the Principal: escalations plus the calibration sample
  results RUN_ID       read the Principal's grades back; refuse while any is missing; only then unblind and apply
                       the pre-registered rule

How each rule of the build brief is held:
  1. Schema. validate_item() checks every field of the D-49 schema (id, set, version, canary, domain, seat,
     task_force, cards[], category, source_kind, drafted_by, drafted_where, checked_by, as_at, prompt, context[],
     reference{answer, must_include[{point, cite}], fatal[]}, objective{type, answer, tolerance}, scoring, holdout),
     with D-46's pairwise scoring and categories A to E. Unknown or missing fields are errors. A private item drafted
     off the node is refused (doc Section 1 L55: no Principal data leaves the node; D-49 VERIFIER rule), and so is a
     D-46 category A or C item drafted off the node. freeze refuses an item nobody has checked.
  2. Privacy of the sets. A set, the eval root, a run, the key and a workbook are refused inside a git working tree
     or a copy of the repository on the node (refuse_repository_path). freeze requires a root-only set and set
     directory, and every later command re-checks the eval root, the run, the frozen set and the key
     (require_private, Run.open). The default root is DEFAULT_EVAL_ROOT, root:root 700.
  3. The engine path. Answers and judges go to the engines through the Engine Arbiter (an EngineGate: the
     orchestrator's own Arbiter over its loopback admin routes in production, OrchestratorGate, the contract
     phase3/loadtest.py uses; an in-process Arbiter in the tests, ArbiterGate) and the engine client
     (engines.LlamaClient, so the per-slot window guard of S52 holds). Engines run one at a time, every task of one
     engine before the next, and the judges start with whatever engine is already resident (engine_order). Nothing
     here imports or calls the router, the chat pipeline (atlas.api), memory, retention, scars, Celery tasks or the
     graph; the system prompt is the production layers 1 to 3 (persona, governance block, domain cards;
     atlas.prompts.PromptBuilder) with no layer 4. The only ledger rows this module writes are `eval-freeze` and
     `eval-unblind` task rows, which carry hashes, counts and the rule, never a prompt or an answer; the Arbiter
     writes its own decision rows for each load (Section 4.2 rule 9), as it does for any load.
  4. Pre-registration. freeze() records the hashes and the rule in the ledger (tasks row, kind `eval-freeze`)
     before any run; run refuses a run whose ledger row is missing (before its first answer) or does not match, and
     results re-checks it. Both refuse a run whose recorded rule differs from the rule this code applies.
  5. Caps and sampling. D-49: max_tokens 16,384 per answer with every truncation (finish "length") recorded. D-46:
     the production caps and sampling per role (D46_ROLES, from deep_think.py and api.py) with each arm's truncation
     rate per category, flagged above 10%.
  6. Judging and the rules: seat_rule() and trial_rule() below, with the constants the recommendation fixed.
  7. The workbook: write_workbook() / read_workbook() (openpyxl, the [eval] extra), in the style of
     tools/build-baseline-workbook.py: yellow response cells with dropdowns. results() unblinds only after every
     escalation and calibration row has the Principal's grade, and every domain (or the trial) whose calibration
     agreement fell below 80% has been graded in full.

What this runner cannot see (stated, not worked around): the Arbiter's single generation slot (Section 4.2 rule 3)
covers requests that pass through the orchestrator. A request this runner sends to an engine's port is one the
orchestrator does not see, as in phase3/loadtest.py, so before each request OrchestratorGate waits until the
orchestrator is not generating and checks that the engine is still resident. Runs belong in a quiet window: the
recommendation budgets 28-59 h (D-49) and 15-35 h (D-46) of exclusive engine time, during which the node serves
nothing else. A run that is interrupted resumes where it stopped.

Markers. Every file:line cited here was read on 2026-10-17 (VERIFIED). INFERRED where said: the roles of D-46
categories D and E, the judges' cap and temperature, and the reading of "the neutral judge agreeing". UNVERIFIED: the
trial engine's API (DwarfStar in a ROCm container, S-51, not built); the runner assumes an OpenAI-compatible chat
endpoint on loopback, which the trial build must confirm.
"""

from __future__ import annotations

import argparse
import contextlib
import datetime as dt
import hashlib
import json
import logging
import math
import os
import random
import re
import shutil
import sqlite3
import stat
import sys
import time
import uuid
from collections.abc import Callable, Collection, Iterator, Mapping, Sequence
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Protocol
from urllib.parse import urlsplit

import httpx

from atlas.arbiter import Arbiter, ArbiterError, read_unit_profile
from atlas.config import (
    ENGINE_KEYS,
    PERSONA_KEYS,
    ConfigError,
    DomainCard,
    EngineSpec,
    Persona,
    Settings,
    load_domain_cards,
    load_engines,
    load_personas,
    parse_env_file,
)
from atlas.deep_think import ADVERSARY_TEMPERATURE, RUBRIC
from atlas.engines import ContextWindowExceeded, EngineError, LlamaClient, SlotWindow, window_for
from atlas.ledger import Ledger
from atlas.prompts import PromptBuilder

log = logging.getLogger("atlas.evalrun")

# --- the sets (recommendations.json D-49 and D-46 "values" and "build_changes_if_approved") -------------------------

SEAT_SET = "seat-test-sets"  # D-49 (T-27); the to-do id of phase3/loadtest.py:325
TRIAL_SET = "trial-task-set"  # D-46 (T-30); the to-do id of phase3/loadtest.py:324
SET_KINDS: tuple[str, ...] = (SEAT_SET, TRIAL_SET)
SCORING: dict[str, str] = {SEAT_SET: "5-3-2+fatal", TRIAL_SET: "pairwise+deep-think-rubric"}
CATEGORIES: dict[str, tuple[str, ...]] = {
    SEAT_SET: ("APPLY", "GROUNDED", "PRECISE", "PRODUCE", "TRAP"),
    TRIAL_SET: ("A", "B", "C", "D", "E"),
}
SOURCE_KINDS: tuple[str, ...] = ("public", "private")
DRAFTED_WHERE: tuple[str, ...] = ("node", "offnode")
# Who may draft an item. The own-draft rule (an engine is never scored on an item it drafted) compares drafted_by with
# the engine keys exactly, so a free-text drafter ("GLM-5.3-Flash", "glm") would void it silently. On the node: an
# engine key (config.ENGINE_KEYS) or the Principal; off the node: the research session only (Claude, not a contestant;
# notes_SETS.md section 5, "Who drafts").
DRAFTERS: dict[str, tuple[str, ...]] = {"node": (*ENGINE_KEYS, "principal"), "offnode": ("offnode-session",)}
# The C0 control characters an .xlsx cell cannot hold (openpyxl's ILLEGAL_CHARACTERS_RE: all but tab, newline and
# carriage return). A form feed copied from a PDF would pass the freeze and then stop the scoring workbook.
CONTROL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
OBJECTIVE_TYPES: tuple[str, ...] = ("exact", "numeric", "flaws", "none")
ITEM_FIELDS: tuple[str, ...] = (
    "id", "set", "version", "canary", "domain", "seat", "task_force", "cards", "category", "source_kind", "drafted_by",
    "drafted_where", "checked_by", "as_at", "prompt", "context", "reference", "objective", "scoring", "holdout",
)
CONTEXT_FIELDS: frozenset[str] = frozenset({"title", "text", "source_url", "compilation"})
REFERENCE_FIELDS: frozenset[str] = frozenset({"answer", "must_include", "fatal"})
OBJECTIVE_FIELDS: frozenset[str] = frozenset({"type", "answer", "tolerance"})
ID_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9._-]{0,63}")
CANARY_MIN_CHARS = 16  # a BIG-bench style canary string (README L144, L236): one per set, long enough to be unique
MAX_CARDS = 3  # Section 8.4 rule 2 (router-rules.json max_domain_cards)
CARD_NUMBERS = range(1, 37)  # Section 8.2: domains 1 to 36
# D-49 "values": wave, then APPLY / GROUNDED / PRECISE / PRODUCE / TRAP per domain (totals 35/15/23/16/11, 100 tasks).
D49_PLAN: dict[str, tuple[int, tuple[int, int, int, int, int]]] = {
    "au-law": (1, (6, 3, 3, 2, 2)),
    "hr": (1, (4, 2, 2, 1, 1)),
    "aviation": (1, (5, 2, 3, 1, 1)),
    "security-osint": (1, (5, 2, 2, 2, 1)),
    "logistics": (2, (3, 1, 4, 1, 1)),
    "aec": (2, (3, 3, 2, 1, 1)),
    "scada": (2, (3, 2, 1, 1, 1)),
    "growth-marketing": (2, (3, 0, 3, 3, 1)),
    "customer-success": (2, (2, 0, 1, 2, 1)),
    "game-design": (2, (1, 0, 2, 2, 1)),
}
D49_PRIVATE_SHARE = 0.30  # "at least 30% from the Principal's own redacted matters"
D49_HOLDOUT_SHARE = 0.20  # "20% held out"
D46_PLAN: dict[str, int] = {"A": 16, "B": 10, "C": 6, "D": 4, "E": 4}
D46_NODE_ONLY: tuple[str, ...] = ("A", "C")  # private TF_OMEGA matters, drafted only on the node (doc L603)
D46_OBJECTIVE: tuple[str, ...] = ("B", "D", "E")  # "objective keys for B, D and E"; "no loss in B, D or E"

# --- engines (CONVENTIONS.md §8 keys; config.ENGINE_KEYS) ------------------------------------------------------------

GLM, DEEPSEEK, GPT_OSS = "glm-5.3-flash", "deepseek-v4-flash", "gpt-oss-120b"
DAY1_ENGINES: tuple[str, ...] = (GLM, DEEPSEEK, GPT_OSS)  # D-49: the three Day 1 seats, three families
SEAT_INCUMBENT = GLM  # Section 6.2: GLM-5.3-Flash carries every D-49 seat, unproven (doc L457)
TRIAL_INCUMBENT = DEEPSEEK  # DeepSeek-V4-Flash-0731, the Apex
# The trial engine's key. The S-51 trial build (not built; doc 15.5 L949) registers it with the Arbiter as the
# exclusive external engine; this runner never starts or stops it, it only checks that it is resident.
TRIAL_KEY = "deepseek-v4.1-flash"
TRIAL_JUDGES: tuple[str, ...] = (GLM, GPT_OSS)  # neither is a DeepSeek model


@dataclass(frozen=True)
class Role:
    """Production sampling and cap for one D-46 category (both arms alike)."""

    name: str
    temperature: float
    top_p: float | None
    max_tokens: int


D49_MAX_TOKENS = 16384  # D-49 VERIFIER: high enough that no reasoning is cut; every truncation is recorded
SCORING_MAX_TOKENS = 1024  # deep_think.py:265, score_trajectories (the adversary's rubric JSON)
SYNTHESIS_MAX_TOKENS = 3072  # deep_think.py:406, run_deep (the deep tier's final synthesis)
# The chat pipeline's cap, under which a TF_OMEGA apex synthesis runs: api.py:183 DEFAULT_MAX_TOKENS, used at
# api.py:1425 when a request sets none (INFERRED that TF_OMEGA takes that path: router.py:15). Copied, not imported:
# this module must not import atlas.api (tests/test_evalrun.py, which also checks that the two agree).
APEX_MAX_TOKENS = 4096
# D-46: "A 1.0 / top-p 1.0 / effort high; B 0.15; C 0.6" (DeepSeek's server default is temperature 1.0, top-p 1.0,
# reasoning effort high: config/engines.json L159; adversary 0.15: deep_think.py:73; synthesis 0.6: deep_think.py:466).
# D and E have no role of their own in the recommendation; they run as the Apex runs A (INFERRED).
D46_ROLES: dict[str, Role] = {
    "A": Role("apex-synthesis", 1.0, 1.0, APEX_MAX_TOKENS),
    "B": Role("adversary", ADVERSARY_TEMPERATURE, None, SCORING_MAX_TOKENS),
    "C": Role("deep-synthesis", 0.6, None, SYNTHESIS_MAX_TOKENS),
    "D": Role("apex-long-dossier", 1.0, 1.0, APEX_MAX_TOKENS),
    "E": Role("apex-computation", 1.0, 1.0, APEX_MAX_TOKENS),
}
TRUNCATION_FLAG_RATE = 0.10  # D-46 VERIFIER: above this, re-run the category at a higher cap (a separate finding)
# Judges grade at the adversary's low-not-zero temperature (Section 9.1: reasoning models loop at exactly zero); the
# cap covers high-effort reasoning before the JSON (INFERRED; the verifier estimated 600-1,500 tokens per call).
JUDGE_TEMPERATURE = ADVERSARY_TEMPERATURE
JUDGE_MAX_TOKENS = 4096
JUDGE_CONTEXT_CHARS = 24000  # longer supplied context (the D-46 32K dossiers) is left out of the judge's prompt

# --- the rules (pre-registered; recorded in the ledger at freeze) ----------------------------------------------------

USABLE_MIN = 7  # usable = 7 or more of 10 with no fatal flag
ESCALATE_GAP = 3  # judges 3 or more points apart -> the Principal grades it
CALIBRATION_FRACTION = 0.10  # D-49: 10% of answers
CALIBRATION_MIN_PER_DOMAIN = 2  # D-49: drawn per domain, as agreement is judged per domain; at least 2 answers each
AGREEMENT_MIN = 0.80  # below this, the domain (D-49) or the trial (D-46) is graded by hand in full (Zheng 2023)
SEAT_KEEP_RATE = 0.80  # GLM keeps a seat at 80% or more usable ...
SEAT_KEEP_MAX_FATAL = 1  # ... with at most 1 fatal item
SEAT_MOVE_LEAD = 3  # another engine takes it only with 3 or more usable items more, the neutral judge agreeing
TRIAL_TASKS = 40
TRIAL_MIN_WINS = 26  # of 40 decided pairs


def binom_tail(k: int, n: int) -> float:
    """P(X >= k) for X ~ Binomial(n, 1/2): the one-sided sign test (ties excluded)."""
    return sum(math.comb(n, i) for i in range(max(k, 0), n + 1)) / 2**n


# The approved threshold is "26 of 40 decided wins (or the table threshold for the number decided)" (D-46 build
# change). The table is the one-sided sign test at alpha 0.05 (notes_SETS.md section 6, :119-125): 30 decided need 20
# (p = 0.049), 40 need 26 (p = 0.040), 50 need 32 (p = 0.033). So alpha is the table's 0.05, not p(26 of 40).
TRIAL_ALPHA = 0.05
TRIAL_MAX_TIE_FRACTION = 0.25  # ties at or below 25%; above, add 10 tasks rather than decide
TRIAL_EXTEND = 10
TRIAL_PROMISE = 0.45  # T-32 "promise": at least 45% preference with no fatal excess (VERIFIER, changed from 55%)
TRIAL_CALIBRATION_TASKS = 10

# --- where things live on the node (CONVENTIONS.md §2 style) ---------------------------------------------------------

DEFAULT_EVAL_ROOT = Path("/srv/atlas/data/eval")  # root:root 700: sets/, runs/<run>/, keys/<run>.key.json
# Copies of the repository on the node: scripts/day1 (atlas-day1.sh) and the package with its venv (02-orchestrator.sh).
REPO_COPIES: tuple[Path, ...] = (Path("/opt/atlas"),)
DIR_MODE, FILE_MODE, FROZEN_MODE = 0o700, 0o600, 0o400
CLIENT_TIMEOUT_S = 900.0  # per streamed chunk; a 16k-token answer at 7 tok/s takes about 40 minutes in all
LOAD_TIMEOUT_S = 1860.0  # api.py's load wait (900 s) plus llama-server@.service TimeoutStartSec (900 s) and margin
QUEUED_RETRY_S = 15.0  # phase3/loadtest.py QUEUED_RETRY_S
IDLE_WAIT_S = 600.0  # phase3/loadtest.py GENERATION_IDLE_WAIT_S
ADMIN_TOKEN_HEADER = "X-Atlas-Token"  # api.py require_admin


class EvalError(RuntimeError):
    """The runner refuses: a bad set, a path in the repository, a missing grade, a ledger record that does not match."""


# --- 1. the schema ---------------------------------------------------------------------------------------------------


def _text(v: Any) -> bool:
    return isinstance(v, str) and bool(v.strip())


def _number(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def _check_context(ctx: Any, where: str) -> list[str]:
    if not isinstance(ctx, list):
        return [f"{where}: context must be a list"]
    errs: list[str] = []
    for i, c in enumerate(ctx):
        w = f"{where}: context[{i}]"
        if not isinstance(c, dict):
            errs.append(f"{w} must be an object")
            continue
        extra = sorted(set(c) - CONTEXT_FIELDS)
        if extra:
            errs.append(f"{w} has unknown field(s) {extra} (allowed: {sorted(CONTEXT_FIELDS)})")
        if not (_text(c.get("text")) or _text(c.get("source_url"))):
            errs.append(f"{w} needs a text or a source_url")
        errs += [f"{w}.{k} must be a string" for k in sorted(CONTEXT_FIELDS & set(c))
                 if c[k] is not None and not isinstance(c[k], str)]
    return errs


def _check_reference(ref: Any, kind: str, where: str) -> list[str]:
    if not isinstance(ref, dict) or set(ref) != REFERENCE_FIELDS:
        return [f"{where}: reference must be an object with exactly {sorted(REFERENCE_FIELDS)}"]
    errs: list[str] = []
    if not _text(ref["answer"]):
        errs.append(f"{where}: reference.answer must be a non-empty string")
    must = ref["must_include"]
    if not isinstance(must, list) or any(not isinstance(m, dict) or set(m) != {"point", "cite"}
                                         or not _text(m["point"]) or not _text(m["cite"]) for m in must):
        errs.append(f"{where}: reference.must_include must be a list of {{point, cite}} with both non-empty "
                    "(pinpoint citations: section and compilation number)")
    elif kind == SEAT_SET and not must:
        errs.append(f"{where}: reference.must_include is empty; a seat item names at least one point")
    if not isinstance(ref["fatal"], list) or not all(_text(f) for f in ref["fatal"]):
        errs.append(f"{where}: reference.fatal must be a list of non-empty strings")
    return errs


def _check_objective(obj: Any, where: str) -> list[str]:
    if not isinstance(obj, dict) or set(obj) != OBJECTIVE_FIELDS:
        return [f"{where}: objective must be an object with exactly {sorted(OBJECTIVE_FIELDS)}"]
    t, answer, tol = obj["type"], obj["answer"], obj["tolerance"]
    if t not in OBJECTIVE_TYPES:
        return [f"{where}: objective.type {t!r} is not one of {OBJECTIVE_TYPES}"]
    if t == "none":
        ok = answer is None and tol is None
        return [] if ok else [f"{where}: objective type none takes answer null and tolerance null"]
    if t == "exact":
        ok = _text(answer) or (isinstance(answer, list) and bool(answer) and all(_text(a) for a in answer))
        if not ok or tol is not None:
            return [f"{where}: objective exact takes a non-empty string (or list of strings) and tolerance null"]
        return []
    if t == "numeric":
        tol_ok = tol is None or (_number(tol) and tol >= 0) or (
            isinstance(tol, str) and re.fullmatch(r"\d+(\.\d+)?%", tol) is not None)
        if not _number(answer) or not tol_ok:
            return [f"{where}: objective numeric takes a number and a tolerance (number >= 0, 'N%' or null)"]
        return []
    ok = isinstance(answer, list) and bool(answer) and all(
        isinstance(f, dict) and set(f) == {"id", "keywords"} and _text(f["id"]) and isinstance(f["keywords"], list)
        and bool(f["keywords"]) and all(_text(k) for k in f["keywords"]) for f in answer)
    if not ok or tol is not None:
        return [f"{where}: objective flaws takes a list of {{id, keywords[]}} and tolerance null"]
    return []


def _workbook_texts(raw: Mapping[str, Any]) -> Iterator[tuple[str, str]]:
    """(field name, text) for every string of an item that the scoring workbook shows: the prompt, the context and the
    reference. Malformed structures are skipped here; the other checks report them."""
    if isinstance(raw.get("prompt"), str):
        yield "prompt", raw["prompt"]
    for i, c in enumerate(raw["context"] if isinstance(raw.get("context"), list) else []):
        for k, v in (c.items() if isinstance(c, dict) else ()):
            if isinstance(v, str):
                yield f"context[{i}].{k}", v
    ref = raw.get("reference") if isinstance(raw.get("reference"), dict) else {}
    if isinstance(ref.get("answer"), str):
        yield "reference.answer", ref["answer"]
    for i, m in enumerate(ref["must_include"] if isinstance(ref.get("must_include"), list) else []):
        for k in ("point", "cite"):
            if isinstance(m, dict) and isinstance(m.get(k), str):
                yield f"reference.must_include[{i}].{k}", m[k]
    for i, f in enumerate(ref["fatal"] if isinstance(ref.get("fatal"), list) else []):
        if isinstance(f, str):
            yield f"reference.fatal[{i}]", f


def validate_item(raw: Any, where: str) -> list[str]:
    """Every reason one JSONL item breaks the schema (empty list: valid). Rule 1."""
    if not isinstance(raw, dict):
        return [f"{where}: not a JSON object"]
    errs: list[str] = []
    unknown, missing = sorted(set(raw) - set(ITEM_FIELDS)), [k for k in ITEM_FIELDS if k not in raw]
    if unknown:
        errs.append(f"{where}: unknown field(s) {unknown}")
    if missing:
        return [*errs, f"{where}: missing field(s) {missing}"]
    kind = raw["set"]
    if kind not in SET_KINDS:
        return [*errs, f"{where}: set {kind!r} is not one of {SET_KINDS}"]
    where = f"{where} ({raw['id']})" if isinstance(raw["id"], str) else where
    if not isinstance(raw["id"], str) or not ID_RE.fullmatch(raw["id"]):
        errs.append(f"{where}: id must match {ID_RE.pattern}")
    if not _text(raw["version"]):
        errs.append(f"{where}: version must be a non-empty string")
    if not isinstance(raw["canary"], str) or len(raw["canary"].strip()) < CANARY_MIN_CHARS:
        errs.append(f"{where}: canary must be a string of at least {CANARY_MIN_CHARS} characters")
    if not _text(raw["domain"]):
        errs.append(f"{where}: domain must be a non-empty string")
    elif kind == SEAT_SET and raw["domain"] not in D49_PLAN:
        errs.append(f"{where}: domain {raw['domain']!r} is not a D-49 domain ({sorted(D49_PLAN)})")
    if raw["seat"] not in PERSONA_KEYS:
        errs.append(f"{where}: seat {raw['seat']!r} is not a persona key ({PERSONA_KEYS})")
    if raw["task_force"] is not None and not _text(raw["task_force"]):
        errs.append(f"{where}: task_force must be a string or null")
    cards = raw["cards"]
    if not isinstance(cards, list) or len(cards) > MAX_CARDS or any(
            not isinstance(c, int) or isinstance(c, bool) or c not in CARD_NUMBERS for c in cards):
        errs.append(f"{where}: cards must be a list of at most {MAX_CARDS} domain numbers 1-36 (Section 8.4 rule 2)")
    if raw["category"] not in CATEGORIES[kind]:
        errs.append(f"{where}: category {raw['category']!r} is not one of {CATEGORIES[kind]} for {kind}")
    if raw["source_kind"] not in SOURCE_KINDS:
        errs.append(f"{where}: source_kind must be one of {SOURCE_KINDS}")
    allowed_by = (DRAFTERS.get(raw["drafted_where"]) if isinstance(raw["drafted_where"], str) else None) or (
        *DRAFTERS["node"], *DRAFTERS["offnode"])
    if raw["drafted_by"] not in allowed_by:
        errs.append(f"{where}: drafted_by {raw['drafted_by']!r} must be one of {list(allowed_by)} for drafted_where "
                    f"{raw['drafted_where']!r} (the own-draft rule compares it with the engine keys exactly)")
    if raw["drafted_where"] not in DRAFTED_WHERE:
        errs.append(f"{where}: drafted_where must be one of {DRAFTED_WHERE}")
    elif raw["drafted_where"] == "offnode" and raw["source_kind"] == "private":
        errs.append(f"{where}: a private item drafted off the node is refused: the Principal's briefs, redacted "
                    "examples and private items are drafted on the node only (doc Section 1 L55; D-49 VERIFIER rule)")
    elif kind == TRIAL_SET and raw["category"] in D46_NODE_ONLY and raw["drafted_where"] != "node":
        errs.append(f"{where}: D-46 category {raw['category']} is a Sensitive TF_OMEGA matter (doc L603), drafted "
                    "only on the node")
    if raw["checked_by"] is not None and not _text(raw["checked_by"]):
        errs.append(f"{where}: checked_by must be a non-empty string or null (not yet checked)")
    try:
        dt.date.fromisoformat(raw["as_at"])
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw["as_at"]):
            raise ValueError(raw["as_at"])
    except (TypeError, ValueError):
        errs.append(f"{where}: as_at must be a date YYYY-MM-DD")
    if not _text(raw["prompt"]):
        errs.append(f"{where}: prompt must be a non-empty string")
    errs += [f"{where}: {name} holds a control character (U+{ord(m.group()):04X}) that the scoring workbook cannot "
             "hold; remove it" for name, value in _workbook_texts(raw)
             if (m := CONTROL_CHARS_RE.search(value)) is not None]
    errs += _check_context(raw["context"], where)
    errs += _check_reference(raw["reference"], kind, where)
    obj_errs = _check_objective(raw["objective"], where)
    errs += obj_errs
    if (not obj_errs and kind == TRIAL_SET and raw["category"] in D46_OBJECTIVE
            and raw["objective"]["type"] == "none"):
        errs.append(f"{where}: D-46 category {raw['category']} needs an objective key (type exact, numeric or flaws)")
    if raw["scoring"] != SCORING[kind]:
        errs.append(f"{where}: scoring must be {SCORING[kind]!r} for {kind}")
    if not isinstance(raw["holdout"], bool):
        errs.append(f"{where}: holdout must be true or false")
    return errs


@dataclass
class EvalSet:
    """A validated set file: its bytes' sha256 and its items in file order."""

    path: Path
    sha256: str
    items: list[dict[str, Any]]
    kind: str
    version: str
    canary: str

    @property
    def by_id(self) -> dict[str, dict[str, Any]]:
        return {it["id"]: it for it in self.items}


def load_set(path: Path, *, require_checked: bool = False) -> EvalSet:
    """Parse and validate a JSONL set (rule 1). One set, one version and one canary per file; ids unique. With
    `require_checked` (freeze), an item nobody has checked is an error. Raises EvalError naming every problem."""
    data = path.read_bytes()
    try:
        lines = data.decode("utf-8").splitlines()
    except UnicodeDecodeError as exc:
        raise EvalError(f"{path}: not UTF-8 ({exc})") from exc
    errs: list[str] = []
    items: list[dict[str, Any]] = []
    for n, line in enumerate(lines, 1):
        if not line.strip():
            continue
        try:
            raw = json.loads(line)
        except json.JSONDecodeError as exc:
            errs.append(f"line {n}: not JSON ({exc.msg})")
            continue
        item_errs = validate_item(raw, f"line {n}")
        errs += item_errs
        if not item_errs:
            items.append(raw)
    if not items and not errs:
        errs.append("the file holds no items")
    for name in ("set", "version", "canary"):
        values = sorted({str(it[name]) for it in items})
        if len(values) > 1:
            errs.append(f"one {name} per file, found {len(values)}: {values[:3]}")
    seen: set[str] = set()
    for it in items:
        if it["id"] in seen:
            errs.append(f"id {it['id']!r} is used twice")
        seen.add(it["id"])
    if require_checked:
        errs += [f"{it['id']}: checked_by is null; the Principal approves each item before the set is frozen"
                 for it in items if it["checked_by"] is None]
    if errs:
        more = f"\n... and {len(errs) - 40} more" if len(errs) > 40 else ""
        raise EvalError(f"{path}: {len(errs)} problem(s):\n" + "\n".join(errs[:40]) + more)
    return EvalSet(path=path, sha256=hashlib.sha256(data).hexdigest(), items=items, kind=items[0]["set"],
                   version=items[0]["version"], canary=items[0]["canary"])


def plan_report(evset: EvalSet) -> list[str]:
    """Advisory lines: the set's counts against the approved values. Never an error: the Principal may drop or halve
    a domain (D-49 "what_would_change_it"), and a wave file holds only its own domains."""
    out: list[str] = []
    items = evset.items
    if evset.kind == SEAT_SET:
        cats = CATEGORIES[SEAT_SET]
        for domain in sorted({it["domain"] for it in items}):
            got = tuple(sum(1 for it in items if it["domain"] == domain and it["category"] == c) for c in cats)
            wave, want = D49_PLAN[domain]
            if got != want:
                out.append(f"{domain} (wave {wave}): {dict(zip(cats, got, strict=True))} against the approved "
                           f"{dict(zip(cats, want, strict=True))}")
            if not got[-1]:
                out.append(f"{domain}: no TRAP item (each domain has at least one)")
            n = sum(got)
            held = sum(1 for it in items if it["domain"] == domain and it["holdout"])
            if held < math.floor(D49_HOLDOUT_SHARE * n):
                out.append(f"{domain}: {held} of {n} held out (the plan holds out 20%)")
        private = sum(1 for it in items if it["source_kind"] == "private")
        if private < D49_PRIVATE_SHARE * len(items):
            out.append(f"{private} of {len(items)} items are private (the plan asks at least 30% from the "
                       "Principal's own redacted matters)")
    else:
        for cat, want in D46_PLAN.items():
            got = sum(1 for it in items if it["category"] == cat)
            if got != want:
                out.append(f"category {cat}: {got} items against the approved {want}")
        public_only = [it["id"] for it in items if it["category"] in D46_OBJECTIVE and it["source_kind"] != "public"]
        if public_only:
            out.append(f"{len(public_only)} B/D/E item(s) are not public (B, D and E come from public material)")
    unchecked = sum(1 for it in items if it["checked_by"] is None)
    if unchecked:
        out.append(f"{unchecked} item(s) not yet checked by the Principal; freeze refuses them")
    return out


# --- 2. private paths ------------------------------------------------------------------------------------------------


def git_tree_of(path: Path) -> Path | None:
    """The git working tree `path` sits in (a `.git` entry in it or a parent), or None."""
    p = path.resolve()
    for d in (p, *p.parents):
        if (d / ".git").exists():
            return d
    return None


def refuse_repository_path(path: Path, what: str) -> None:
    """Rule 2: sets, keys, answers and workbooks never live in a repository or a copy of one."""
    p = path.resolve()
    tree = git_tree_of(p)
    if tree is not None:
        raise EvalError(f"{what} {p} is inside the git working tree {tree}: evaluation sets, keys, answers and "
                        "workbooks are private and never live in a repository (D-49; doc Section 1 L55)")
    for root in (*REPO_COPIES, Path(__file__).resolve().parent):
        if p == root or root in p.parents:
            raise EvalError(f"{what} {p} is inside {root}, a copy of the repository on the node (CONVENTIONS.md §2)")


def require_private(path: Path, owner_uid: int, what: str) -> None:
    """Rule 2: owned by `owner_uid` (root on the node) with no group or other bits; never a symlink."""
    st = path.lstat()
    if stat.S_ISLNK(st.st_mode):
        raise EvalError(f"{what} {path} is a symlink; give the real path")
    if st.st_uid != owner_uid or stat.S_IMODE(st.st_mode) & 0o077:
        want = "700" if stat.S_ISDIR(st.st_mode) else "600 or 400"
        raise EvalError(f"{what} {path} must be root-only: owner uid {owner_uid}, mode {want} (found uid {st.st_uid}, "
                        f"mode {stat.S_IMODE(st.st_mode):o}); chown and chmod it, the sets are private")


def ensure_private_dir(path: Path, owner_uid: int, what: str) -> Path:
    refuse_repository_path(path, what)
    if not path.exists():
        path.mkdir(mode=DIR_MODE, parents=True)
        os.chmod(path, DIR_MODE)  # mkdir's mode is masked by the umask
    require_private(path, owner_uid, what)
    return path


def write_private(path: Path, data: bytes, mode: int = FILE_MODE) -> None:
    """Write `data` to `path` whole (a temporary file, then rename) with `mode` from the first byte."""
    tmp = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    fd = os.open(tmp, os.O_WRONLY | os.O_CREAT | os.O_EXCL, FILE_MODE)
    try:
        os.write(fd, data)
        os.fsync(fd)
    finally:
        os.close(fd)
    os.chmod(tmp, mode)
    os.replace(tmp, path)


def append_jsonl(path: Path, record: Mapping[str, Any]) -> None:
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_APPEND, FILE_MODE)
    try:
        os.write(fd, (json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n").encode("utf-8"))
        os.fsync(fd)
    finally:
        os.close(fd)


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def canonical(obj: Any) -> bytes:
    return json.dumps(obj, sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode("utf-8")


# --- 4. the ledger (pre-registration) --------------------------------------------------------------------------------


def _ledger_owner(db: Path, euid: int | None) -> tuple[int, int] | None:
    """(uid, gid) to write the ledger as, or None for this process. Root writing a ledger the service owns would leave
    root-owned -wal/-shm files the service cannot write (admin.py _root_owner_conflict), so root writes as the owner."""
    euid = os.geteuid() if euid is None else euid
    if euid != 0 or not db.exists():
        return None
    st = db.stat()
    return None if st.st_uid == 0 else (st.st_uid, st.st_gid)


def with_ledger(db: Path, fn: Callable[[Ledger], Any], *, euid: int | None = None) -> Any:
    """Run `fn(ledger)` on ATLAS_DB_PATH and return its JSON-able result; as the database's owner when this process is
    root and the owner is not (a forked child drops to that uid and gid first)."""

    def body() -> Any:
        ledger = Ledger(db)
        try:
            ledger.init_db()
            return fn(ledger)
        finally:
            ledger.close()

    owner = _ledger_owner(db, euid)
    if owner is None:
        return body()
    rfd, wfd = os.pipe()
    pid = os.fork()
    if pid == 0:  # pragma: no cover - the child's lines run in another process
        os.close(rfd)
        code = 1
        try:
            os.setgroups([])
            os.setgid(owner[1])
            os.setuid(owner[0])
            with os.fdopen(wfd, "wb") as fh:
                fh.write(json.dumps(body()).encode("utf-8"))
            code = 0
        except BaseException as exc:
            sys.stderr.write(f"atlas-admin eval: ledger as uid {owner[0]}: {type(exc).__name__}: {exc}\n")
        finally:
            os._exit(code)
    os.close(wfd)
    with os.fdopen(rfd, "rb") as fh:
        data = fh.read()
    _, status = os.waitpid(pid, 0)
    if os.waitstatus_to_exitcode(status) != 0:
        raise EvalError(f"the ledger call as uid {owner[0]} (the owner of {db}) failed; see the line above")
    return json.loads(data)


def ledger_record(db: Path, kind: str, payload: Mapping[str, Any], *, euid: int | None = None) -> str:
    """One `tasks` row (status done) carrying `payload`: hashes, counts and rules only. Returns its id."""
    task_id = f"eval-{uuid.uuid4().hex}"
    with_ledger(db, lambda led: led.insert_task(kind, task_id=task_id, status="done", payload=dict(payload)),
                euid=euid)
    return task_id


def ledger_row(db: Path, task_id: str, *, euid: int | None = None) -> dict[str, Any] | None:
    return with_ledger(db, lambda led: led.get_task(task_id), euid=euid)


# --- the run directory ------------------------------------------------------------------------------------------------


def preregistered_rule(kind: str) -> dict[str, Any]:
    """The decision rule fixed before any run (rule 4); its sha256 goes to the ledger with the set's and the key's."""
    if kind == SEAT_SET:
        return {
            "set": SEAT_SET, "scale": "substance 0-5, sources 0-3, clarity 0-2, plus FATAL", "usable_min": USABLE_MIN,
            "judges": "the two Day 1 engines that did not write the answer", "escalate_gap": ESCALATE_GAP,
            "escalate_on_fatal": True, "calibration_fraction": CALIBRATION_FRACTION, "agreement_min": AGREEMENT_MIN,
            "calibration": "per domain: ceil(calibration_fraction x the domain's answers), at least "
                           "calibration_min_per_domain, from its non-escalated answers, seeded from the key and the "
                           "domain", "calibration_min_per_domain": CALIBRATION_MIN_PER_DOMAIN,
            "incumbent": SEAT_INCUMBENT, "keep_rate": SEAT_KEEP_RATE, "keep_max_fatal": SEAT_KEEP_MAX_FATAL,
            "incumbent_at_bar_keeps": True,
            "move_lead": SEAT_MOVE_LEAD, "move_needs_neutral_judge_lead": SEAT_MOVE_LEAD,
            "no_engine_at_bar": "mandatory verifier pass and a raised approval tier",
            "own_drafts_excluded": True, "max_tokens": D49_MAX_TOKENS, "judge_max_tokens": JUDGE_MAX_TOKENS,
        }
    return {
        "set": TRIAL_SET, "incumbent": TRIAL_INCUMBENT, "trial": TRIAL_KEY, "judges": list(TRIAL_JUDGES),
        "orders_per_judge": 2, "min_wins": TRIAL_MIN_WINS, "of_decided": TRIAL_TASKS, "alpha": round(TRIAL_ALPHA, 6),
        "threshold": "smallest k with P(X >= k | decided, 1/2) <= alpha", "max_tie_fraction": TRIAL_MAX_TIE_FRACTION,
        "extend_by": TRIAL_EXTEND, "fatal": "trial fatal errors <= incumbent fatal errors",
        "no_loss_categories": list(D46_OBJECTIVE), "d47_floor_required": True, "promise_preference": TRIAL_PROMISE,
        "calibration_tasks": TRIAL_CALIBRATION_TASKS, "agreement_min": AGREEMENT_MIN,
        "roles": {c: asdict(r) for c, r in D46_ROLES.items()}, "truncation_flag_rate": TRUNCATION_FLAG_RATE,
        "judge_max_tokens": JUDGE_MAX_TOKENS,
    }


_LABEL_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"


def make_key(evset: EvalSet, contestants: Sequence[str], *, run_id: str, rng: random.Random) -> dict[str, Any]:
    """The blinding key: an opaque label per (task, contestant), and for D-46 a random X/Y order per task. Committed
    (its sha256 in the ledger) before any answer exists."""
    used: set[str] = set()

    def label(prefix: str) -> str:
        while True:
            s = prefix + "-" + "".join(rng.choice(_LABEL_ALPHABET) for _ in range(6))
            if s not in used:
                used.add(s)
                return s

    answers: dict[str, dict[str, str]] = {}
    pairs: dict[str, dict[str, str]] = {}
    for it in evset.items:
        for engine in contestants:
            answers[label("A")] = {"task": it["id"], "engine": engine}
        if evset.kind == TRIAL_SET:
            arms = list(contestants)
            rng.shuffle(arms)
            pairs[it["id"]] = {"label": label("P"), "x": arms[0], "y": arms[1]}
    seed = "".join(rng.choice("0123456789abcdef") for _ in range(64))
    return {"run_id": run_id, "set_sha256": evset.sha256, "seed": seed, "answers": answers, "pairs": pairs}


@dataclass
class Run:
    """One frozen run: <eval_root>/runs/<run_id>/ and its key, re-verified against the manifest on every open."""

    root: Path
    run_id: str
    dir: Path
    manifest: dict[str, Any]
    evset: EvalSet
    key: dict[str, Any]
    _labels: dict[tuple[str, str], str] = field(default_factory=dict, repr=False)

    @property
    def kind(self) -> str:
        return self.evset.kind

    @property
    def answers_path(self) -> Path:
        return self.dir / "answers.jsonl"

    @property
    def judgements_path(self) -> Path:
        return self.dir / "judgements.jsonl"

    @property
    def runlog_path(self) -> Path:
        return self.dir / "runlog.jsonl"

    @property
    def contestants(self) -> list[str]:
        return list(self.manifest["contestants"])

    def label_of(self, task: str, engine: str) -> str:
        if not self._labels:
            self._labels = {(v["task"], v["engine"]): k for k, v in self.key["answers"].items()}
        return self._labels[(task, engine)]

    def writer(self, label: str) -> str:
        return str(self.key["answers"][label]["engine"])

    @classmethod
    def open(cls, eval_root: Path, run_id: str, *, owner_uid: int) -> Run:
        if not ID_RE.fullmatch(run_id):
            raise EvalError(f"run id {run_id!r} must match {ID_RE.pattern}")
        root = eval_root
        run_dir = root / "runs" / run_id
        if not run_dir.is_dir():
            raise EvalError(f"no run {run_id} under {root / 'runs'} (freeze the set first)")
        for p, what in ((root, "eval root"), (run_dir, "run directory")):
            refuse_repository_path(p, what)
            require_private(p, owner_uid, what)
        manifest = json.loads((run_dir / "manifest.json").read_text(encoding="utf-8"))
        set_path = run_dir / "set.jsonl"
        key_path = Path(manifest["key_path"])
        for p, what in ((set_path, "frozen set"), (key_path, "blinding key")):
            refuse_repository_path(p, what)
            require_private(p, owner_uid, what)
        evset = load_set(set_path, require_checked=True)
        if evset.sha256 != manifest["set_sha256"]:
            raise EvalError(f"{set_path} has changed since it was frozen (sha256 {evset.sha256} against "
                            f"{manifest['set_sha256']}); a frozen set is never edited")
        key_bytes = key_path.read_bytes()
        if sha256_bytes(key_bytes) != manifest["key_sha256"]:
            raise EvalError(f"{key_path} does not match the key recorded at freeze; the labels are void")
        return cls(root=root, run_id=run_id, dir=run_dir, manifest=manifest, evset=evset,
                   key=json.loads(key_bytes))

    def item(self, task: str) -> dict[str, Any]:
        return self.evset.by_id[task]

    def log_event(self, event: Mapping[str, Any]) -> None:
        append_jsonl(self.runlog_path, {"ts": time.time(), **event})


def freeze(set_path: Path, *, eval_root: Path, db: Path, owner_uid: int, run_id: str | None = None,
           rng: random.Random | None = None, euid: int | None = None) -> Run:
    """Rules 2 and 4: check the set is private and approved, copy it into a new run (0400), write the blinding key
    (0600, outside the repository) and record the hashes and the pre-registered rule in the ledger. Nothing has run
    before this returns; on a failed ledger write the run and the key are removed again."""
    refuse_repository_path(set_path, "set")
    require_private(set_path.parent, owner_uid, "set directory")
    require_private(set_path, owner_uid, "set")
    evset = load_set(set_path, require_checked=True)
    rng = rng or random.SystemRandom()
    stamp = dt.datetime.now(dt.UTC).strftime("%Y%m%d")
    tag = "d49" if evset.kind == SEAT_SET else "d46"
    run_id = run_id or re.sub(r"[^A-Za-z0-9._-]", "-", f"{tag}-{evset.version}-{stamp}-{uuid.uuid4().hex[:4]}")[:64]
    if not ID_RE.fullmatch(run_id):
        raise EvalError(f"run id {run_id!r} must match {ID_RE.pattern}")
    root = ensure_private_dir(eval_root, owner_uid, "eval root")
    runs = ensure_private_dir(root / "runs", owner_uid, "runs directory")
    keys = ensure_private_dir(root / "keys", owner_uid, "keys directory")
    run_dir = runs / run_id
    key_path = keys / f"{run_id}.key.json"
    if run_dir.exists() or key_path.exists():
        raise EvalError(f"run {run_id} already exists; a frozen run is never re-frozen (choose another --run-id)")
    contestants = list(DAY1_ENGINES) if evset.kind == SEAT_SET else [TRIAL_INCUMBENT, TRIAL_KEY]
    judges = list(DAY1_ENGINES) if evset.kind == SEAT_SET else list(TRIAL_JUDGES)
    key = make_key(evset, contestants, run_id=run_id, rng=rng)
    key_bytes = json.dumps(key, sort_keys=True, indent=1).encode("utf-8")
    rule = preregistered_rule(evset.kind)
    run_dir.mkdir(mode=DIR_MODE)
    os.chmod(run_dir, DIR_MODE)
    try:
        write_private(run_dir / "set.jsonl", set_path.read_bytes(), FROZEN_MODE)
        write_private(key_path, key_bytes, FILE_MODE)
        manifest: dict[str, Any] = {
            "run_id": run_id, "set": evset.kind, "version": evset.version, "items": len(evset.items),
            "set_sha256": evset.sha256, "key_sha256": sha256_bytes(key_bytes), "key_path": str(key_path),
            "rule": rule, "rule_sha256": sha256_bytes(canonical(rule)), "contestants": contestants,
            "judges": judges, "frozen_at": dt.datetime.now(dt.UTC).isoformat(timespec="seconds"),
            "source": str(set_path.resolve()),
        }
        payload = {k: manifest[k] for k in ("run_id", "set", "version", "items", "set_sha256", "key_sha256", "rule",
                                            "rule_sha256", "contestants", "judges")}
        manifest["ledger_task_id"] = ledger_record(db, "eval-freeze", payload, euid=euid)
        write_private(run_dir / "manifest.json", json.dumps(manifest, indent=1, sort_keys=True).encode("utf-8"))
    except BaseException:
        shutil.rmtree(run_dir, ignore_errors=True)
        key_path.unlink(missing_ok=True)
        raise
    return Run.open(root, run_id, owner_uid=owner_uid)


def verify_preregistration(run: Run, db: Path, *, required: bool, euid: int | None = None) -> str:
    """Rule 4: the `eval-freeze` row must exist and carry this run's hashes and rule. `required` (run): a missing row
    refuses. Otherwise (results, possibly weeks later) a row the 30-day retention archived (D9, ledger.RETENTION_S) is
    said so and the manifest's hashes, checked by Run.open, stand. Either way the rule this code applies must be the
    rule pre-registered at freeze: results() and the runner use the module constants, so a constant changed after the
    freeze would otherwise decide a frozen run under another rule. Returns "ledger" or "archived"."""
    if sha256_bytes(canonical(run.manifest["rule"])) != run.manifest["rule_sha256"]:
        raise EvalError(f"run {run.run_id}: the manifest's rule no longer matches its recorded sha256")
    now = preregistered_rule(run.kind)
    if canonical(now) != canonical(run.manifest["rule"]):
        changed = sorted(k for k in set(now) | set(run.manifest["rule"])
                         if canonical(now.get(k)) != canonical(run.manifest["rule"].get(k)))
        raise EvalError(f"run {run.run_id}: the rule in this code differs from the rule pre-registered for it at "
                        f"freeze ({', '.join(changed)}); decide it with the code that froze it, or re-freeze")
    tid = run.manifest.get("ledger_task_id")
    row = ledger_row(db, tid, euid=euid) if tid else None
    if row is None:
        if required:
            raise EvalError(f"run {run.run_id}: no eval-freeze row {tid} in the ledger {db}; the hashes and the rule "
                            "are recorded before any run (re-freeze)")
        log.warning("run %s: the eval-freeze row %s is no longer in the ledger (archived after 30 days, D9); the "
                    "manifest's hashes are used", run.run_id, tid)
        return "archived"
    payload = json.loads(row.get("payload_json") or "{}")
    for k in ("set_sha256", "key_sha256", "rule_sha256"):
        if payload.get(k) != run.manifest.get(k):
            raise EvalError(f"run {run.run_id}: the ledger records {k} {payload.get(k)} but the run has "
                            f"{run.manifest.get(k)}; refusing")
    return "ledger"


def verify_for_run(run: Run, db: Path, *, euid: int | None = None) -> str:
    """verify_preregistration for `run`: the ledger row is required only before the first answer exists. A D-49 run
    takes 28-59 h over several quiet nights, and the 30-day retention (D9) may archive the eval-freeze row before its
    answers or judges finish; a run that already passed with the row then resumes on the manifest's hashes, as
    results() does, rather than being re-frozen and losing every answer."""
    return verify_preregistration(run, db, required=not run.answers_path.exists(), euid=euid)


# --- 3. the engine path -----------------------------------------------------------------------------------------------


class EngineGate(Protocol):
    """Loads go through the Engine Arbiter; a generation waits for its single generation slot."""

    def load(self, key: str, *, task_id: str) -> None: ...

    def is_resident(self, key: str) -> bool: ...

    def generation(self, key: str, *, task_id: str) -> contextlib.AbstractContextManager[None]: ...


def _serving(status: Mapping[str, Any], key: str) -> bool:
    return any(r.get("engine") == key and not r.get("loading") and not r.get("unloading")
               for r in status.get("resident") or [])


class ArbiterGate:
    """An Arbiter this process owns (the tests drive it with StubController and StubProbe). In production the Arbiter
    is the orchestrator's (Section 4.2: one Arbiter), reached through OrchestratorGate."""

    def __init__(self, arbiter: Arbiter, *, load_wait_s: float = LOAD_TIMEOUT_S) -> None:
        self.arbiter = arbiter
        self.load_wait_s = load_wait_s

    def load(self, key: str, *, task_id: str) -> None:
        dec = self.arbiter.request_load(key, task_id=task_id, wait_s=self.load_wait_s)
        if not dec.granted:
            raise EvalError(f"the Arbiter answered {dec.decision.value} for {key}: {dec.reason}")

    def is_resident(self, key: str) -> bool:
        return _serving(self.arbiter.status(), key)

    def generation(self, key: str, *, task_id: str) -> contextlib.AbstractContextManager[None]:
        return self.arbiter.acquire_generation(key, task_id=task_id, timeout_s=IDLE_WAIT_S)


class OrchestratorGate:
    """The orchestrator's Arbiter over its loopback admin routes (the contract phase3/loadtest.py uses): POST
    /arbiter/load, GET /arbiter/status, and GET /health for the generation slot ("generating", "generation_queue").
    Never a systemctl call: Section 4.2 says every load passes through the Arbiter."""

    def __init__(self, url: str, *, admin_token: str | None = None, transport: httpx.BaseTransport | None = None,
                 sleep: Callable[[float], None] = time.sleep, clock: Callable[[], float] = time.monotonic) -> None:
        parts = urlsplit(url)
        if parts.scheme != "http" or (parts.hostname or "").lower() not in LlamaClient.LOOPBACK_HOSTS:
            raise EvalError(f"the orchestrator is reached on loopback only (http://127.0.0.1:<ORCH_PORT>); refusing "
                            f"{url!r}")
        headers = {ADMIN_TOKEN_HEADER: admin_token} if admin_token else {}
        self._http = httpx.Client(base_url=url.rstrip("/"), timeout=LOAD_TIMEOUT_S, headers=headers,
                                  trust_env=False, transport=transport)
        self._sleep, self._clock = sleep, clock

    def _call(self, method: str, path: str, body: Mapping[str, Any] | None = None) -> tuple[int, Any]:
        try:
            r = self._http.request(method, path, json=body)
        except httpx.HTTPError as exc:
            raise EvalError(f"{method} {path}: {type(exc).__name__}: {exc} (is atlas-orchestrator serving?)") from exc
        try:
            return r.status_code, r.json()
        except ValueError:
            return r.status_code, r.text[:300]

    def load(self, key: str, *, task_id: str) -> None:
        deadline = self._clock() + LOAD_TIMEOUT_S
        while True:
            code, body = self._call("POST", "/arbiter/load", {"engine": key, "task_id": task_id})
            if code != 200 or not isinstance(body, dict):
                raise EvalError(f"POST /arbiter/load {key} answered {code}: {str(body)[:300]}")
            if body.get("granted"):
                return
            if body.get("decision") == "queued" and self._clock() < deadline:
                self._sleep(QUEUED_RETRY_S)
                continue
            raise EvalError(f"the Arbiter answered {body.get('decision')} for {key}: {body.get('reason')}")

    def is_resident(self, key: str) -> bool:
        code, body = self._call("GET", "/arbiter/status")
        if code != 200 or not isinstance(body, dict):
            raise EvalError(f"GET /arbiter/status answered {code}: {str(body)[:300]}")
        return _serving(body, key)

    @contextlib.contextmanager
    def generation(self, key: str, *, task_id: str) -> Iterator[None]:
        deadline = self._clock() + IDLE_WAIT_S
        while True:
            code, body = self._call("GET", "/health")
            if code != 200 or not isinstance(body, dict):
                raise EvalError(f"GET /health answered {code} while checking the generation slot")
            if not body.get("generating") and not body.get("generation_queue"):
                break
            if self._clock() > deadline:
                raise EvalError(f"the orchestrator is still generating after {IDLE_WAIT_S:.0f}s: something uses the "
                                "assistant during the evaluation (Section 4.2 rule 3); stop it and resume the run")
            self._sleep(5.0)
        if not self.is_resident(key):
            raise EvalError(f"{key} is no longer resident (another request evicted it?); resume the run in a quiet "
                            "window")
        yield


def read_admin_token(path: str | Path | None) -> str | None:
    """ORCH_ADMIN_TOKEN_FILE: `ORCH_ADMIN_TOKEN=...` or the bare token (phase3/loadtest.py's reader); never logged."""
    if not path:
        return None
    for raw in Path(path).read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if line and not line.startswith("#"):
            return line.split("=", 1)[1].strip().strip("'\"") if line.startswith("ORCH_ADMIN_TOKEN=") else line
    raise EvalError(f"{path} is empty")


def client_factory(engines: Mapping[str, EngineSpec], *, env_dir: Path | None,
                   external: Mapping[str, tuple[str, int]] | None = None,
                   transport_for: Callable[[str], httpx.BaseTransport | None] | None = None,
                   ) -> Callable[[str], LlamaClient]:
    """An engine client per key, with the per-slot window of the profile the unit runs (as atlas.api.engine_window
    reads it: the env file, else engines.json). `external` maps a key the trial build runs (TRIAL_KEY) to its
    loopback URL and window."""

    def make(key: str) -> LlamaClient:
        transport = transport_for(key) if transport_for else None
        if external and key in external:
            url, tokens = external[key]
            return LlamaClient(url, timeout_s=CLIENT_TIMEOUT_S, transport=transport,
                               window=SlotWindow(key, tokens, "--trial-window (the trial build's profile)"))
        if key not in engines:
            raise EvalError(f"{key} is not an engines.json key")
        spec = engines[key]
        prof = read_unit_profile(env_dir / f"{key}.env") if env_dir is not None else None
        window = window_for(spec, prof.ctx, prof.parallel, source=prof.describe()) if prof else window_for(spec)
        return LlamaClient(spec.base_url, timeout_s=CLIENT_TIMEOUT_S, transport=transport, window=window)

    return make


def engine_order(engines: Sequence[str], resident: str | None) -> list[str]:
    """Group by engine to minimise swaps: the engine already resident first, the rest in their given order."""
    order = list(engines)
    if resident in order:
        order.remove(resident)
        order.insert(0, resident)
    return order


def _resident_of(gate: EngineGate, candidates: Sequence[str]) -> str | None:
    return next((k for k in candidates if gate.is_resident(k)), None)


# --- prompts ----------------------------------------------------------------------------------------------------------


def production_temperature(persona: Persona) -> float | None:
    """The chat pipeline's rule (api.py:1417-1421): a numeric sampling_temperature as given, a phrase such as
    arthur.md's "low-not-zero" as 0.15, none as the engine's server default."""
    t = persona.sampling_temperature
    if isinstance(t, (int, float)) and not isinstance(t, bool):
        return float(t)
    return ADVERSARY_TEMPERATURE if isinstance(t, str) else None


class SeatPrompts:
    """The production system prompt, layers 1 to 3 of Section 4.4 (persona core, governance block, the item's domain
    cards; atlas.prompts.PromptBuilder) and NO layer 4: no memory and no scars are read; the task carries its own
    context (D-49 build changes: "read only the task's supplied context")."""

    def __init__(self, personas: Mapping[str, Persona], cards: Mapping[int, DomainCard]) -> None:
        self.personas = personas
        self.cards = cards
        self._builder = PromptBuilder(personas)

    def problems(self, items: Sequence[Mapping[str, Any]]) -> list[str]:
        out = [f"{it['id']}: seat {it['seat']!r} has no persona file" for it in items
               if it["seat"] not in self.personas]
        out += [f"{it['id']}: domain card {n} is not in config/domains/cards" for it in items for n in it["cards"]
                if n not in self.cards]
        return out

    def system(self, item: Mapping[str, Any]) -> str:
        cards = [self.cards[n] for n in item["cards"]]
        return self._builder.build_system_prompt(item["seat"], cards, None).text


@dataclass(frozen=True)
class Sampling:
    temperature: float | None
    top_p: float | None
    max_tokens: int
    role: str


def sampling_for(kind: str, item: Mapping[str, Any], prompts: SeatPrompts) -> Sampling:
    """Rule 5: D-49 the seat's own sampling at 16,384 tokens; D-46 the production role of the category."""
    if kind == SEAT_SET:
        return Sampling(production_temperature(prompts.personas[item["seat"]]), None, D49_MAX_TOKENS, "seat")
    role = D46_ROLES[item["category"]]
    return Sampling(role.temperature, role.top_p, role.max_tokens, role.name)


def task_message(item: Mapping[str, Any]) -> str:
    """The user turn: the prompt, then the supplied context, numbered, with its source and compilation."""
    parts = [str(item["prompt"]).strip()]
    if item["context"]:
        parts.append("Supplied context (cite it by number and source):")
        for i, c in enumerate(item["context"], 1):
            head = " | ".join(str(c[k]) for k in ("title", "source_url", "compilation") if c.get(k))
            parts.append(f"[{i}] {head}".rstrip() + (f"\n{str(c['text']).strip()}" if c.get("text") else ""))
    return "\n\n".join(parts)


# --- blinding ---------------------------------------------------------------------------------------------------------

_THINK = re.compile(r"<think>.*?(?:</think>|\Z)", re.S | re.I)
_HARMONY_ANALYSIS = re.compile(r"<\|channel\|>\s*analysis\s*<\|message\|>.*?(?:<\|end\|>|\Z)", re.S)
_SPECIAL = re.compile(r"<\|[^|>\s]{1,40}\|>")
# Self-identification: the contestants' and their vendors' names (and the commonest others) become "[model]". A task
# that itself names a vendor loses that word in the blinded copy only; the answer file keeps the original.
_SELF_ID = re.compile(
    r"\b(?:Chat)?GLM(?:-[\w.]+)?\b|\bZhipu(?:\s?AI)?\b|\bZ\.ai\b|\bDeepSeek(?:-[\w.]+)?\b|\bgpt-oss(?:-\w+)?\b|"
    r"\bGPT-?\d[\w.]*\b|\bChatGPT\b|\bOpenAI\b|\bQwen[\w.-]*\b|\bClaude\b|\bAnthropic\b|\bLlama[\w.-]*\b|"
    r"\bMistral\b|\bGemma\b|\bGemini\b",
    re.I,
)


def blind(text: str) -> str:
    """The text a judge and the Principal see: reasoning blocks stripped, self-identification removed, the markdown
    normalised (one bullet glyph, '1.' numbering, no trailing spaces, at most one blank line)."""
    t = _THINK.sub("", text)
    t = _HARMONY_ANALYSIS.sub("", t)
    t = _SPECIAL.sub("", t)
    t = _SELF_ID.sub("[model]", t)
    t = re.sub(r"(?m)^(\s*)[*+•]\s+", r"\1- ", t)
    t = re.sub(r"(?m)^(\s*)(\d+)\)\s+", r"\1\2. ", t)
    t = re.sub(r"(?m)[ \t]+$", "", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def blinded_answer(rec: Mapping[str, Any]) -> str:
    if rec.get("error"):
        return "[no answer: the engine could not take this task]"
    return blind(str(rec.get("text") or "")) or "[empty answer]"


# --- 3 and 5. answers -------------------------------------------------------------------------------------------------


def _answer_one(client: LlamaClient, gate: EngineGate, engine: str, item: Mapping[str, Any], label: str,
                sampling: Sampling, system: str, task_id: str) -> dict[str, Any]:
    messages = [{"role": "system", "content": system}, {"role": "user", "content": task_message(item)}]
    rec: dict[str, Any] = {"label": label, "task": item["id"], "requested_max_tokens": sampling.max_tokens,
                           "role": sampling.role}
    extra = {"top_p": sampling.top_p} if sampling.top_p is not None else {}
    started = time.monotonic()
    with gate.generation(engine, task_id=task_id):
        try:
            res = client.chat(messages, stream=True, max_tokens=sampling.max_tokens, temperature=sampling.temperature,
                              **extra)
        except ContextWindowExceeded as exc:
            rec.update(text="", error=f"context window: {exc}", finish_reason=None, truncated=False, max_tokens=None)
            return rec
    fit = client.last_fit
    sent = fit.max_tokens if fit is not None else sampling.max_tokens
    timings = res.timings
    rec.update(text=res.text, error=None, finish_reason=res.finish_reason, truncated=res.finish_reason == "length",
               max_tokens=sent, clamped=sent < sampling.max_tokens, seconds=round(time.monotonic() - started, 1),
               predicted_n=timings.predicted_n if timings else None,
               decode_tok_s=round(timings.predicted_per_second, 2) if timings else None)
    return rec


def truncation_summary(run: Run, records: Sequence[Mapping[str, Any]]) -> dict[str, dict[str, Any]]:
    """Per engine: answers, truncations (finish "length"), clamps and errors, overall and per category; a D-46 arm
    whose rate in a category passes 10% is flagged (rule 5)."""
    out: dict[str, dict[str, Any]] = {}
    for rec in records:
        engine = run.writer(rec["label"])
        cat = run.item(rec["task"])["category"]
        e = out.setdefault(engine, {"answers": 0, "truncated": 0, "clamped": 0, "errors": 0, "by_category": {}})
        c = e["by_category"].setdefault(cat, {"answers": 0, "truncated": 0})
        e["answers"] += 1
        c["answers"] += 1
        if rec.get("truncated"):
            e["truncated"] += 1
            c["truncated"] += 1
        e["clamped"] += 1 if rec.get("clamped") else 0
        e["errors"] += 1 if rec.get("error") else 0
    for e in out.values():
        e["rate"] = round(e["truncated"] / e["answers"], 4) if e["answers"] else 0.0
        e["flagged"] = sorted(cat for cat, c in e["by_category"].items()
                              if run.kind == TRIAL_SET and c["answers"] and c["truncated"] / c["answers"]
                              > TRUNCATION_FLAG_RATE)
    return out


def run_answers(run: Run, gate: EngineGate, client_for: Callable[[str], LlamaClient], prompts: SeatPrompts, *,
                engines: Collection[str] | None = None, external: Collection[str] = (),
                say: Callable[[str], None] = print) -> dict[str, dict[str, Any]]:
    """Every task on one engine, then the next (rules 3 and 5). Resumable: an answer already in answers.jsonl is kept.
    An engine in `external` (the trial arm) is never loaded here: it must already be resident."""
    wanted = run.contestants if engines is None else [e for e in run.contestants if e in engines]
    unknown = sorted(set(engines or ()) - set(run.contestants))
    if unknown:
        raise EvalError(f"{unknown} are not contestants of run {run.run_id} ({run.contestants})")
    problems = prompts.problems(run.evset.items)
    if problems:
        raise EvalError("the config cannot build these prompts:\n" + "\n".join(problems))
    done = {r["label"] for r in read_jsonl(run.answers_path)}
    for engine in engine_order(wanted, _resident_of(gate, wanted)):
        todo = [it for it in run.evset.items if run.label_of(it["id"], engine) not in done]
        if not todo:
            continue
        task_id = f"eval-{run.run_id}"
        if engine in external:
            if not gate.is_resident(engine):
                raise EvalError(f"{engine} is not resident: the trial build (S-51) starts it and registers it with the "
                                "Arbiter as the exclusive external engine before its arm runs (doc 15.5)")
        else:
            gate.load(engine, task_id=task_id)
        say(f"{engine}: {len(todo)} answer(s) to write")
        run.log_event({"event": "answers-start", "engine": engine, "todo": len(todo)})
        client = client_for(engine)
        try:
            for it in todo:
                label = run.label_of(it["id"], engine)
                rec = _answer_one(client, gate, engine, it, label, sampling_for(run.kind, it, prompts),
                                  prompts.system(it), f"{task_id}-{label}")
                append_jsonl(run.answers_path, rec)
                done.add(label)
        finally:
            client.close()
        mine = [r for r in read_jsonl(run.answers_path) if run.writer(r["label"]) == engine]
        summary = truncation_summary(run, mine).get(engine, {})
        run.log_event({"event": "answers-done", "engine": engine, "summary": summary})
        say(f"{engine}: {summary.get('answers', 0)} answers, {summary.get('truncated', 0)} truncated "
            f"({100 * summary.get('rate', 0.0):.1f}%), {summary.get('errors', 0)} could not be sent"
            + (f"; ABOVE 10% in {summary['flagged']}: re-run at a higher cap, a separate finding"
               if summary.get("flagged") else ""))
    return truncation_summary(run, read_jsonl(run.answers_path))


# --- 6. judging -------------------------------------------------------------------------------------------------------

SEAT_JUDGE_SYSTEM = (
    "You are a strict, neutral grader for the A.T.L.A.S. seat tests. You did not write the answer and you do not know "
    "who did. Grade it against the reference, which the Principal has checked. Marks: substance 0 to 5 (correct, and "
    "covers every must-include point); sources 0 to 3 (the authorities and citations are right and pinpointed; one "
    "invented citation makes sources 0); clarity 0 to 2 (clear, direct and usable as given). FATAL is true when the "
    "answer invents an authority or citation, applies foreign law as Australian, gives unsafe or unlawful advice, "
    "states a wrong figure that changes the decision, or does anything on the item's fatal list. Reply with one JSON "
    'object and nothing else: {"substance": 0-5, "sources": 0-3, "clarity": 0-2, "fatal": true or false, '
    '"fatal_reason": "..."}'
)
PAIR_JUDGE_SYSTEM = (
    "You are a strict, neutral judge. You compare two answers to the same task; you do not know who wrote either, "
    "and their order is random. Weigh these criteria equally: "
    + "; ".join(f"{name} ({question})" for name, question in RUBRIC)
    + ". Use the reference, which the Principal has checked. An answer has a FATAL error when it invents a fact, an "
    "authority or a figure, states a wrong figure that changes the decision, or does anything on the item's fatal "
    'list. Reply with one JSON object and nothing else: {"better": "1" or "2" or "tie", "fatal_1": true or false, '
    '"fatal_2": true or false, "reason": "..."}'
)


def _reference_block(item: Mapping[str, Any]) -> str:
    ref, obj = item["reference"], item["objective"]
    lines = [f"Answer: {ref['answer']}"]
    if ref["must_include"]:
        lines.append("Must include:")
        lines += [f"- {m['point']} [{m['cite']}]" for m in ref["must_include"]]
    if ref["fatal"]:
        lines.append("Fatal if:")
        lines += [f"- {f}" for f in ref["fatal"]]
    if obj["type"] == "flaws":
        lines.append("Planted flaws: " + "; ".join(f"{f['id']} ({', '.join(f['keywords'])})" for f in obj["answer"]))
    elif obj["type"] != "none":
        tol = f" (tolerance {obj['tolerance']})" if obj["tolerance"] is not None else ""
        lines.append(f"Key ({obj['type']}): {obj['answer']}{tol}")
    return "\n".join(lines)


def _context_for_judge(item: Mapping[str, Any]) -> str:
    body = task_message({**item, "prompt": ""}).strip()
    if not body:
        return "(none)"
    if len(body) > JUDGE_CONTEXT_CHARS:
        return (f"(The supplied context, {len(body):,} characters, is left out so the grader's window holds the "
                "answers; grade against the reference.)")
    return body


def seat_judge_messages(item: Mapping[str, Any], answer: str) -> list[dict[str, str]]:
    user = (f"=== Task ===\n{item['prompt']}\n=== Supplied context ===\n{_context_for_judge(item)}\n"
            f"=== Reference (checked by the Principal) ===\n{_reference_block(item)}\n"
            f"=== Answer to grade ===\n{answer}\n=== End ===")
    return [{"role": "system", "content": SEAT_JUDGE_SYSTEM}, {"role": "user", "content": user}]


def pair_judge_messages(item: Mapping[str, Any], first: str, second: str) -> list[dict[str, str]]:
    user = (f"=== Task ===\n{item['prompt']}\n=== Supplied context ===\n{_context_for_judge(item)}\n"
            f"=== Reference (checked by the Principal) ===\n{_reference_block(item)}\n"
            f"=== Answer 1 ===\n{first}\n=== Answer 2 ===\n{second}\n=== End ===")
    return [{"role": "system", "content": PAIR_JUDGE_SYSTEM}, {"role": "user", "content": user}]


_JSON_OBJ = re.compile(r"\{[^{}]*\}", re.S)


def _last_json(text: str) -> dict[str, Any] | None:
    for m in reversed(list(_JSON_OBJ.finditer(blind(text)))):
        try:
            obj = json.loads(m.group(0))
        except json.JSONDecodeError:
            continue
        if isinstance(obj, dict):
            return obj
    return None


def _int_in(v: Any, lo: int, hi: int) -> int | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, str) and re.fullmatch(r"\s*\d+\s*", v):
        v = int(v)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return v if isinstance(v, int) and lo <= v <= hi else None


def _flag(v: Any) -> bool | None:
    if isinstance(v, bool):
        return v
    if isinstance(v, str) and v.strip().lower() in ("yes", "true", "no", "false"):
        return v.strip().lower() in ("yes", "true")
    return None


@dataclass(frozen=True)
class Grade:
    """One D-49 grade, a judge's or the Principal's (the Allens scale plus FATAL)."""

    substance: int
    sources: int
    clarity: int
    fatal: bool

    @property
    def total(self) -> int:
        return self.substance + self.sources + self.clarity

    @property
    def usable(self) -> bool:
        return self.total >= USABLE_MIN and not self.fatal


def make_grade(substance: Any, sources: Any, clarity: Any, fatal: Any) -> Grade | None:
    s, so, c, f = _int_in(substance, 0, 5), _int_in(sources, 0, 3), _int_in(clarity, 0, 2), _flag(fatal)
    return None if None in (s, so, c, f) else Grade(s, so, c, f)  # type: ignore[arg-type]


def parse_grade(text: str) -> Grade | None:
    obj = _last_json(text)
    return None if obj is None else make_grade(obj.get("substance"), obj.get("sources"), obj.get("clarity"),
                                               obj.get("fatal"))


def parse_pair(text: str) -> tuple[str, bool, bool] | None:
    """('1' | '2' | 'tie', fatal in answer 1, fatal in answer 2), or None when the reply does not parse."""
    obj = _last_json(text)
    if obj is None:
        return None
    better = str(obj.get("better", "")).strip().lower()
    f1, f2 = _flag(obj.get("fatal_1")), _flag(obj.get("fatal_2"))
    if better not in ("1", "2", "tie") or f1 is None or f2 is None:
        return None
    return better, f1, f2


def _judge_call(client: LlamaClient, gate: EngineGate, judge: str, messages: list[dict[str, str]],
                parse: Callable[[str], Any], task_id: str) -> tuple[Any, dict[str, Any]]:
    """One judge call with one literal retry when the reply does not parse (as deep_think.score_trajectories)."""
    meta: dict[str, Any] = {"truncated": False, "error": None, "attempts": 0}
    for _ in range(2):
        meta["attempts"] += 1
        with gate.generation(judge, task_id=task_id):
            try:
                res = client.chat(messages, stream=True, max_tokens=JUDGE_MAX_TOKENS, temperature=JUDGE_TEMPERATURE)
            except ContextWindowExceeded as exc:
                meta["error"] = f"context window: {exc}"
                return None, meta
        meta["truncated"] = res.finish_reason == "length"
        parsed = parse(res.text)
        if parsed is not None:
            return parsed, meta
    return None, meta


def run_judges(run: Run, gate: EngineGate, client_for: Callable[[str], LlamaClient], *,
               say: Callable[[str], None] = print) -> None:
    """D-49: each answer graded by the two Day 1 engines that did not write it. D-46: each pair judged by GLM and
    gpt-oss in both orders (four verdicts per task). One judge engine at a time, the resident one first."""
    answers = {r["label"]: r for r in read_jsonl(run.answers_path)}
    missing = sorted(set(run.key["answers"]) - set(answers))
    if missing:
        raise EvalError(f"{len(missing)} answer(s) are not written yet (run --phase answers for every contestant "
                        "first); judging starts when the set is complete")
    done = read_jsonl(run.judgements_path)
    judges = list(run.manifest["judges"])
    task_id = f"eval-{run.run_id}-judge"
    for judge in engine_order(judges, _resident_of(gate, judges)):
        if run.kind == SEAT_SET:
            have = {r["label"] for r in done if r["judge"] == judge}
            todo: list[Any] = [label for label in sorted(answers, key=lambda lb: (answers[lb]["task"], lb))
                               if run.writer(label) != judge and label not in have]
        else:
            have = {(r["task"], r["order"]) for r in done if r["judge"] == judge}
            todo = [(it["id"], order) for it in run.evset.items for order in ("xy", "yx")
                    if (it["id"], order) not in have]
        if not todo:
            continue
        gate.load(judge, task_id=task_id)
        say(f"{judge}: {len(todo)} judgement(s) to make")
        client = client_for(judge)
        try:
            for entry in todo:
                rec = (_seat_judgement(run, client, gate, judge, entry, answers, task_id) if run.kind == SEAT_SET
                       else _pair_judgement(run, client, gate, judge, entry, answers, task_id))
                append_jsonl(run.judgements_path, rec)
        finally:
            client.close()
        run.log_event({"event": "judges-done", "judge": judge, "judgements": len(todo)})


def _seat_judgement(run: Run, client: LlamaClient, gate: EngineGate, judge: str, label: str,
                    answers: Mapping[str, Mapping[str, Any]], task_id: str) -> dict[str, Any]:
    item = run.item(answers[label]["task"])
    grade, meta = _judge_call(client, gate, judge, seat_judge_messages(item, blinded_answer(answers[label])),
                              parse_grade, f"{task_id}-{label}")
    return {"label": label, "judge": judge, "grade": asdict(grade) if grade else None, **meta}


def _pair_judgement(run: Run, client: LlamaClient, gate: EngineGate, judge: str, entry: tuple[str, str],
                    answers: Mapping[str, Mapping[str, Any]], task_id: str) -> dict[str, Any]:
    task, order = entry
    pair = run.key["pairs"][task]
    slots = ("x", "y") if order == "xy" else ("y", "x")
    texts = [blinded_answer(answers[run.label_of(task, pair[s])]) for s in slots]
    parsed, meta = _judge_call(client, gate, judge, pair_judge_messages(run.item(task), texts[0], texts[1]),
                               parse_pair, f"{task_id}-{pair['label']}-{order}")
    rec: dict[str, Any] = {"task": task, "judge": judge, "order": order, "winner": None, "fatal_x": None,
                           "fatal_y": None, **meta}
    if parsed is not None:
        better, f1, f2 = parsed
        rec["winner"] = "tie" if better == "tie" else slots[int(better) - 1]
        rec[f"fatal_{slots[0]}"], rec[f"fatal_{slots[1]}"] = f1, f2
    return rec


# --- 6. aggregation, escalation and calibration -----------------------------------------------------------------------


@dataclass
class SeatAnswer:
    """One D-49 answer with its two judges' grades, whether it goes to the Principal, and the Principal's grade."""

    label: str
    task: str
    engine: str
    domain: str
    judges: dict[str, Grade | None]
    escalate: bool
    reasons: list[str]
    principal: Grade | None = None

    @property
    def judge_mean(self) -> float | None:
        grades = list(self.judges.values())
        if not grades or any(g is None for g in grades):
            return None
        return sum(g.total for g in grades if g is not None) / len(grades)

    @property
    def judge_fatal(self) -> bool:
        return any(g is not None and g.fatal for g in self.judges.values())

    @property
    def judges_usable(self) -> bool:
        mean = self.judge_mean
        return mean is not None and mean >= USABLE_MIN and not self.judge_fatal

    @property
    def usable(self) -> bool:
        return self.principal.usable if self.principal is not None else self.judges_usable

    @property
    def fatal(self) -> bool:
        return self.principal.fatal if self.principal is not None else self.judge_fatal


def aggregate_seat(run: Run) -> dict[str, SeatAnswer]:
    """Every answer with its two judges; escalated when a reply did not parse, the judges are 3 or more points apart,
    or either flags FATAL (rule 6)."""
    by_label: dict[str, dict[str, Grade | None]] = {}
    for r in read_jsonl(run.judgements_path):
        g = r.get("grade")
        by_label.setdefault(r["label"], {})[r["judge"]] = Grade(**g) if g else None
    out: dict[str, SeatAnswer] = {}
    for label, meta in run.key["answers"].items():
        expected = [j for j in run.manifest["judges"] if j != meta["engine"]]
        got = by_label.get(label, {})
        if sorted(got) != sorted(expected):
            raise EvalError(f"{label}: judgements from {sorted(got)}, expected {expected}; finish --phase judges")
        reasons: list[str] = []
        grades = [got[j] for j in expected]
        if any(g is None for g in grades):
            reasons.append("a judge's reply did not parse")
        elif abs(grades[0].total - grades[1].total) >= ESCALATE_GAP:  # type: ignore[union-attr]
            reasons.append(f"judges {ESCALATE_GAP} or more points apart")
        reasons += [f"{j} flagged FATAL" for j in expected if got[j] is not None and got[j].fatal]  # type: ignore
        out[label] = SeatAnswer(label, meta["task"], meta["engine"], run.item(meta["task"])["domain"],
                                {j: got[j] for j in expected}, bool(reasons), reasons)
    return out


@dataclass
class PairTask:
    """One D-46 task: four verdicts (two judges, both orders) mapped to the slots X and Y of the key."""

    task: str
    label: str
    category: str
    x: str
    y: str
    verdicts: dict[tuple[str, str], str | None]
    fatal: dict[str, bool]
    judge_winner: str | None  # "x" | "y" | "tie" when both judges agree in both orders, else None (split)
    escalate: bool
    reasons: list[str]
    principal: tuple[str, bool, bool] | None = None  # (winner slot, fatal X, fatal Y)

    @property
    def winner(self) -> str | None:
        return self.principal[0] if self.principal is not None else self.judge_winner

    def arm_fatal(self, arm: str) -> bool:
        slot = "x" if arm == self.x else "y"
        if self.principal is not None:
            return self.principal[1] if slot == "x" else self.principal[2]
        return self.fatal[slot]

    def winner_arm(self) -> str | None:
        w = self.winner
        return {"x": self.x, "y": self.y}.get(w) if w in ("x", "y") else w


def aggregate_pairs(run: Run) -> dict[str, PairTask]:
    """A judge whose two orders disagree, a reply that did not parse, or two judges that disagree make a split; a
    split counts as a tie until the Principal grades it, and is escalated with every FATAL flag (rule 6)."""
    verdicts: dict[str, dict[tuple[str, str], dict[str, Any]]] = {}
    for r in read_jsonl(run.judgements_path):
        verdicts.setdefault(r["task"], {})[(r["judge"], r["order"])] = r
    out: dict[str, PairTask] = {}
    judges = list(run.manifest["judges"])
    for it in run.evset.items:
        task, pair = it["id"], run.key["pairs"][it["id"]]
        got = verdicts.get(task, {})
        expected = [(j, o) for j in judges for o in ("xy", "yx")]
        if sorted(got) != sorted(expected):
            raise EvalError(f"{task}: {len(got)} of 4 verdicts; finish --phase judges")
        reasons: list[str] = []
        per_judge = []
        for j in judges:
            a, b = got[(j, "xy")]["winner"], got[(j, "yx")]["winner"]
            if a is None or b is None:
                reasons.append(f"{j}: a reply did not parse")
            elif a != b:
                reasons.append(f"{j}: the two orders disagree")
            per_judge.append(a if a == b else None)
        winner = per_judge[0] if per_judge and all(v is not None and v == per_judge[0] for v in per_judge) else None
        if winner is None and not reasons:
            reasons.append("the judges disagree")
        fatal = {s: any(bool(r.get(f"fatal_{s}")) for r in got.values()) for s in ("x", "y")}
        reasons += [f"FATAL flagged in {s.upper()}" for s in ("x", "y") if fatal[s]]
        out[task] = PairTask(task, pair["label"], it["category"], pair["x"], pair["y"],
                             {k: v["winner"] for k, v in got.items()}, fatal, winner, bool(reasons), reasons)
    return out


def calibration_sample(run: Run, pool: Sequence[str], size: int, salt: str = "") -> list[str]:
    """A random sample seeded from the blinding key (and `salt`: a D-49 domain), so it is fixed before any grade is
    seen."""
    rng = random.Random(f"{run.key['seed']}:calibration" + (f":{salt}" if salt else ""))
    pool = sorted(pool)
    return sorted(rng.sample(pool, min(size, len(pool))))


def principal_rows(run: Run, *, all_rows: bool = False, domains: Collection[str] = ()) -> dict[str, str]:
    """Row key -> why the Principal grades it: "escalation", "calibration", or "full" (a domain, or the whole
    trial, graded by hand because calibration agreement fell below 80%)."""
    rows: dict[str, str] = {}
    if run.kind == SEAT_SET:
        seat = aggregate_seat(run)
        rows.update({a.label: "escalation" for a in seat.values() if a.escalate})
        # Per domain, as agreement is judged per domain: a small domain (6 tasks, 18 answers) drawn from the whole
        # pool could get no calibration row at all, and one row makes agreement 0% or 100%.
        for domain in sorted({a.domain for a in seat.values()}):
            mine = [a for a in seat.values() if a.domain == domain]
            size = max(CALIBRATION_MIN_PER_DOMAIN, math.ceil(CALIBRATION_FRACTION * len(mine)))
            rows.update({lb: "calibration" for lb in calibration_sample(
                run, [a.label for a in mine if not a.escalate], size, salt=domain)})
        rows.update({a.label: "full" for a in seat.values()
                     if (all_rows or a.domain in domains) and a.label not in rows})
    else:
        pairs = aggregate_pairs(run)
        rows.update({p.label: "escalation" for p in pairs.values() if p.escalate})
        rows.update({lb: "calibration" for lb in calibration_sample(run, [p.label for p in pairs.values()
                                                                          if not p.escalate],
                                                                    TRIAL_CALIBRATION_TASKS)})
        rows.update({p.label: "full" for p in pairs.values() if all_rows and p.label not in rows})
    return rows


# --- 6. the rules -----------------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class Tally:
    n: int
    usable: int
    fatal: int

    @property
    def rate(self) -> float:
        return self.usable / self.n if self.n else 0.0


def meets_bar(t: Tally) -> bool:
    """D-49 keep bar: 80% or more usable with at most 1 fatal item."""
    return t.n > 0 and t.rate >= SEAT_KEEP_RATE and t.fatal <= SEAT_KEEP_MAX_FATAL


def seat_rule(tallies: Mapping[str, Tally], leads: Mapping[str, tuple[int, int]],
              incumbent: str = SEAT_INCUMBENT) -> dict[str, Any]:
    """D-49 seat rule, in the approved order: "GLM keeps a seat at 80% or more usable with at most 1 fatal item;
    another engine takes it only if it leads by 3 or more usable items with the neutral judge agreeing". The
    incumbent that meets the bar keeps the seat, whatever another engine's lead. Only when it misses the bar may
    another engine take it. `leads[x]` = (x's usable items minus the incumbent's on the items neither drafted, by the
    final grades; the same lead by the neutral judge's grades alone, the engine that is neither). An engine takes the
    seat only when both leads are 3 or more (reading the neutral judge's agreement as the same 3-item lead is
    INFERRED). When no engine reaches the bar the domain gets a mandatory verifier pass and a raised approval tier."""
    movers = sorted(((lead, x) for x, (lead, neutral) in leads.items()
                     if lead >= SEAT_MOVE_LEAD and neutral >= SEAT_MOVE_LEAD), reverse=True)
    bar = {e: meets_bar(t) for e, t in tallies.items()}
    if bar.get(incumbent):
        seat, decision = incumbent, "keeps"
    elif len(movers) > 1 and movers[0][0] == movers[1][0]:
        seat, decision = incumbent, f"contested: {movers[0][1]} and {movers[1][1]} lead equally; the Principal decides"
    elif movers:
        seat, decision = movers[0][1], "moves"
    else:
        seat, decision = incumbent, "keeps, unproven"
    verifier = not any(bar.values())
    return {"seat": seat, "decision": decision, "meets_bar": bar, "verifier_pass_required": verifier,
            "raised_approval_tier": verifier}


def wins_needed(decided: int, alpha: float = TRIAL_ALPHA) -> int | None:
    """The fewest wins of `decided` pairs whose one-sided sign-test p is at or below alpha 0.05, the approved table's
    threshold (20 of 30, 26 of 40, 32 of 50); None when no count can reach it (fewer than 5 decided pairs)."""
    for k in range(decided + 1):
        if binom_tail(k, decided) <= alpha + 1e-12:
            return k
    return None


def trial_rule(*, wins: int, losses: int, ties: int, fatal_trial: int, fatal_incumbent: int,
               objective: Mapping[str, tuple[float, float]], floor_met: bool | None) -> dict[str, Any]:
    """D-46 promotion rule, fixed before unblinding. `objective[cat]` = (trial score, incumbent score) for B, D, E."""
    total = wins + losses + ties
    decided = wins + losses
    tie_fraction = ties / total if total else 0.0
    k = wins_needed(decided)
    preference = wins / decided if decided else 0.0
    out: dict[str, Any] = {
        "wins": wins, "losses": losses, "ties": ties, "decided": decided, "tie_fraction": round(tie_fraction, 4),
        "wins_needed": k, "fatal_trial": fatal_trial, "fatal_incumbent": fatal_incumbent,
        "objective": {c: list(v) for c, v in objective.items()}, "floor_met": floor_met,
        "preference": round(preference, 4),
        "promise": decided > 0 and preference >= TRIAL_PROMISE and fatal_trial <= fatal_incumbent,
    }
    if tie_fraction > TRIAL_MAX_TIE_FRACTION:
        out.update(decision="extend", promote=False,
                   reason=f"ties {100 * tie_fraction:.1f}% are above 25%: add {TRIAL_EXTEND} tasks (re-blinded) "
                          "before deciding")
        return out
    checks = {
        "wins": k is not None and wins >= k,
        "fatal": fatal_trial <= fatal_incumbent,
        "objective": all(t >= i for t, i in objective.values()),
        "floor": floor_met is True,
    }
    out["checks"] = checks
    if all(checks.values()):
        out.update(decision="promote", promote=True, reason="every condition holds")
    elif floor_met is None and all(v for c, v in checks.items() if c != "floor"):
        out.update(decision="held", promote=False, reason="the D-47 floor is not yet confirmed (T-31)")
    else:
        failed = [c for c, v in checks.items() if not v]
        out.update(decision="keep-0731", promote=False, reason="failed: " + ", ".join(failed))
    return out


# --- objective checks (D-46 B, D, E) ----------------------------------------------------------------------------------

_NUM = re.compile(r"[-+]?\d[\d,]*(?:\.\d+)?")


def _answer_number(text: str) -> float | None:
    lines = [ln for ln in text.splitlines() if re.match(r"\s*answer\s*:", ln, re.I)]
    hits = _NUM.findall(lines[-1]) if lines else _NUM.findall(text)
    if not hits:
        return None
    try:
        return float(hits[0 if lines else -1].replace(",", ""))
    except ValueError:
        return None


def check_objective(obj: Mapping[str, Any], text: str) -> tuple[float, float]:
    """(score, maximum) of one answer against its key: exact facts found, a number within tolerance (the last
    'ANSWER:' line, else the last number), or planted flaws named (any keyword of the flaw)."""
    t = obj["type"]
    if t == "none":
        return 0.0, 0.0
    if t == "exact":
        wants = [obj["answer"]] if isinstance(obj["answer"], str) else list(obj["answer"])
        body = re.sub(r"\s+", " ", text).lower()
        return float(sum(1 for w in wants if re.sub(r"\s+", " ", w).strip().lower() in body)), float(len(wants))
    if t == "numeric":
        x = _answer_number(text)
        tol = obj["tolerance"]
        allowed = (abs(float(obj["answer"])) * float(tol[:-1]) / 100 if isinstance(tol, str)
                   else float(tol or 0.0))
        return (1.0 if x is not None and abs(x - float(obj["answer"])) <= allowed + 1e-12 else 0.0), 1.0
    low = text.lower()
    caught = sum(1 for f in obj["answer"] if any(k.lower() in low for k in f["keywords"]))
    return float(caught), float(len(obj["answer"]))


# --- 7. the workbook --------------------------------------------------------------------------------------------------

HEADER_ROW = 4
CELL_LIMIT = 32000  # Excel holds 32,767 characters per cell; longer text goes to a file beside the workbook
SEAT_SHEET, PAIR_SHEET, META_SHEET, README_SHEET = "Grades", "Pairs", "_meta", "Read me"
SEAT_COLUMNS: tuple[tuple[str, int], ...] = (
    ("Label", 11), ("Domain", 14), ("Type", 11), ("Task", 50), ("Supplied context", 40), ("Reference answer", 45),
    ("Must include (point [citation])", 45), ("Fatal if", 30), ("Answer", 80), ("Substance 0-5", 11),
    ("Sources 0-3", 11), ("Clarity 0-2", 11), ("Fatal", 9), ("Notes", 30),
)
SEAT_INPUTS: dict[str, list[str]] = {
    "Substance 0-5": [str(i) for i in range(6)], "Sources 0-3": [str(i) for i in range(4)],
    "Clarity 0-2": [str(i) for i in range(3)], "Fatal": ["No", "Yes"],
}
PAIR_COLUMNS: tuple[tuple[str, int], ...] = (
    ("Pair", 11), ("Type", 7), ("Task", 50), ("Supplied context", 40), ("Reference", 45), ("Answer X", 75),
    ("Answer Y", 75), ("Verdict", 13), ("Fatal in X", 10), ("Fatal in Y", 10), ("Notes", 30),
)
PAIR_VERDICTS: dict[str, str] = {"X better": "x", "Y better": "y", "Tie": "tie"}
PAIR_INPUTS: dict[str, list[str]] = {"Verdict": list(PAIR_VERDICTS), "Fatal in X": ["No", "Yes"],
                                     "Fatal in Y": ["No", "Yes"]}


def _openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError as exc:  # the [eval] extra; Day 1 does not install it (nothing on Day 1 runs this module)
        raise EvalError("openpyxl is not installed in this venv: install the package's [eval] extra, as root: "
                        "/opt/atlas/venv-uv/bin/uv pip install --python /opt/atlas/venv/bin/python "
                        "-e '/opt/atlas/orchestrator[eval]' (phase2/02-orchestrator.sh's installer)") from exc
    return openpyxl


def _must_text(item: Mapping[str, Any]) -> str:
    return "\n".join(f"- {m['point']} [{m['cite']}]" for m in item["reference"]["must_include"])


def _reference_text(item: Mapping[str, Any]) -> str:
    return _reference_block(item)


def write_workbook(run: Run, out: Path, *, owner_uid: int, all_rows: bool = False,
                   domains: Collection[str] = ()) -> tuple[Path, int]:
    """Rule 7: the blinded workbook: labels only, never an engine name; the reasoning and self-identification
    stripped (blind). Response cells are yellow with dropdowns, as tools/build-baseline-workbook.py builds them.
    Returns (path, rows)."""
    opx = _openpyxl()
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    def safe(value: Any) -> Any:
        """A cell value openpyxl accepts: a control character in an engine's answer (or a context extract) would raise
        IllegalCharacterError, which is not a ValueError, and stop the workbook of a finished run."""
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    refuse_repository_path(out, "workbook")
    require_private(out.parent, owner_uid, "workbook directory")
    rows_for = principal_rows(run, all_rows=all_rows, domains=domains)
    answers = {r["label"]: r for r in read_jsonl(run.answers_path)}
    overflow_dir = out.with_name(out.stem + "-texts")

    def cell_text(text: str, name: str) -> str:
        if len(text) <= CELL_LIMIT:
            return text
        if not overflow_dir.exists():
            overflow_dir.mkdir(mode=DIR_MODE)
            os.chmod(overflow_dir, DIR_MODE)
        write_private(overflow_dir / f"{name}.md", text.encode("utf-8"))
        return text[:CELL_LIMIT - 200] + f"\n[... continued in {overflow_dir.name}/{name}.md beside this workbook]"

    data: list[list[Any]] = []
    if run.kind == SEAT_SET:
        columns, inputs, sheet_name = SEAT_COLUMNS, SEAT_INPUTS, SEAT_SHEET
        seat = aggregate_seat(run)
        for lb in sorted(rows_for, key=lambda x: (seat[x].domain, seat[x].task, x)):
            it = run.item(seat[lb].task)
            ctx = task_message({**it, "prompt": ""}).strip()
            data.append([lb, it["domain"], it["category"], it["prompt"], cell_text(ctx, f"{lb}-context"),
                         it["reference"]["answer"], _must_text(it), "\n".join(it["reference"]["fatal"]),
                         cell_text(blinded_answer(answers[lb]), f"{lb}-answer"), None, None, None, None, None])
    else:
        columns, inputs, sheet_name = PAIR_COLUMNS, PAIR_INPUTS, PAIR_SHEET
        pairs = {p.label: p for p in aggregate_pairs(run).values()}
        for lb in sorted(rows_for, key=lambda x: (pairs[x].category, pairs[x].task)):
            p = pairs[lb]
            it = run.item(p.task)
            ctx = task_message({**it, "prompt": ""}).strip()
            ax = blinded_answer(answers[run.label_of(p.task, p.x)])
            ay = blinded_answer(answers[run.label_of(p.task, p.y)])
            data.append([lb, it["category"], it["prompt"], cell_text(ctx, f"{lb}-context"), _reference_text(it),
                         cell_text(ax, f"{lb}-X"), cell_text(ay, f"{lb}-Y"), None, None, None, None])

    font = Font(name="Arial", size=10)
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3864")
    input_fill = PatternFill("solid", fgColor="FFFF00")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = opx.Workbook()
    readme = wb.active
    readme.title = README_SHEET
    lines = [
        f"A.T.L.A.S. blind scoring: {run.kind} {run.manifest['version']} (run {run.run_id})",
        "",
        "Every answer here is labelled; no engine is named, and the reasoning and any self-identification have been "
        "removed. Grade each row in the yellow cells, then return this file. Nothing is unblinded until every row "
        "has your grade.",
        "",
    ]
    if run.kind == SEAT_SET:
        lines += [
            "Scale (the Allens benchmark's, plus FATAL): substance 0-5, sources 0-3, clarity 0-2.",
            "Fatal: an invented authority or citation, foreign law applied as Australian, unsafe or unlawful advice, "
            "a wrong figure that changes the decision, or anything on the row's 'Fatal if' list.",
            "Usable means 7 or more with no fatal error.",
        ]
    else:
        lines += [
            "Each row is one task with two answers, X and Y, in a random order. Pick the better one, or Tie, and mark "
            "any fatal error in either.",
            "Criteria, equally weighted: " + "; ".join(name for name, _ in RUBRIC) + ".",
        ]
    lines += ["", "Some rows are where the two engine judges disagreed or one flagged a fatal error; others are a "
              "random sample that checks the judges. They are not marked, on purpose."]
    for i, text in enumerate(lines, 1):
        c = readme.cell(row=i, column=1, value=safe(text))
        c.font, c.alignment = (Font(name="Arial", size=14, bold=True) if i == 1 else font), wrap
    readme.column_dimensions["A"].width = 120

    ws = wb.create_sheet(sheet_name)
    ws["A1"] = f"{sheet_name}: {len(data)} row(s) to grade"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = "Yellow cells are yours. Long text continues in the files named in the cell."
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="555555")
    names = [c for c, _ in columns]
    for j, (name, width) in enumerate(columns, 1):
        c = ws.cell(row=HEADER_ROW, column=j, value=name)
        c.font, c.fill, c.alignment, c.border = head_font, head_fill, wrap, border
        ws.column_dimensions[get_column_letter(j)].width = width
    for i, row in enumerate(data, HEADER_ROW + 1):
        for j, value in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=safe(value))
            c.font, c.alignment, c.border = font, wrap, border
            if names[j - 1] in inputs or names[j - 1] == "Notes":
                c.fill = input_fill
    last = HEADER_ROW + max(len(data), 1)
    for name, options in inputs.items():
        col = get_column_letter(names.index(name) + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                            showDropDown=False, promptTitle="Choose", prompt="Pick one: " + " / ".join(options))
        dv.error, dv.errorTitle = "Pick a value from the list.", "Invalid entry"
        dv.showErrorMessage, dv.errorStyle = True, "stop"
        ws.add_data_validation(dv)
        dv.add(f"{col}{HEADER_ROW + 1}:{col}{last}")
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=2)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(columns))}{last}"

    meta = wb.create_sheet(META_SHEET)
    for i, (k, v) in enumerate((("run_id", run.run_id), ("set", run.kind), ("set_sha256", run.manifest["set_sha256"]),
                                ("rows", len(data))), 1):
        meta.cell(row=i, column=1, value=k)
        meta.cell(row=i, column=2, value=v)
    meta.sheet_state = "hidden"

    tmp = out.with_name(f".{out.name}.{os.getpid()}.tmp")
    wb.save(tmp)
    os.chmod(tmp, FILE_MODE)
    os.replace(tmp, out)
    run.log_event({"event": "workbook", "path": str(out), "rows": len(data),
                   "reasons": {r: sum(1 for v in rows_for.values() if v == r) for r in set(rows_for.values())}})
    return out, len(data)


def read_workbook(path: Path, run: Run) -> tuple[dict[str, Any], list[str]]:
    """The Principal's grades from a returned workbook: {row key: Grade | (winner slot, fatal X, fatal Y)} for every
    fully graded row, and the problems (a row half graded, a value outside its list, the wrong run)."""
    opx = _openpyxl()
    wb = opx.load_workbook(path, data_only=True)
    if META_SHEET not in wb.sheetnames:
        raise EvalError(f"{path}: no {META_SHEET} sheet; not a workbook this runner wrote")
    meta = {str(r[0]): r[1] for r in wb[META_SHEET].iter_rows(values_only=True) if r and r[0] is not None}
    if meta.get("run_id") != run.run_id or meta.get("set_sha256") != run.manifest["set_sha256"]:
        raise EvalError(f"{path} belongs to run {meta.get('run_id')}, not {run.run_id}")
    sheet = SEAT_SHEET if run.kind == SEAT_SET else PAIR_SHEET
    ws = wb[sheet]
    header = [c.value for c in ws[HEADER_ROW]]
    col = {name: header.index(name) for name, _ in (SEAT_COLUMNS if run.kind == SEAT_SET else PAIR_COLUMNS)}
    grades: dict[str, Any] = {}
    problems: list[str] = []

    def val(row: Sequence[Any], name: str) -> str:
        v = row[col[name]]
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return "" if v is None else str(v).strip()

    for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        if not row or row[0] is None:
            continue
        key = str(row[0])
        if run.kind == SEAT_SET:
            vals = [val(row, n) for n in SEAT_INPUTS]
            if not any(vals):
                continue
            g = make_grade(*vals[:3], {"yes": True, "no": False}.get(vals[3].lower()))
            if g is None:
                problems.append(f"{key}: substance, sources, clarity and fatal must all be chosen from their lists "
                                f"(got {vals})")
            else:
                grades[key] = g
        else:
            vals = [val(row, n) for n in PAIR_INPUTS]
            if not any(vals):
                continue
            winner = PAIR_VERDICTS.get(vals[0])
            fx, fy = ({"yes": True, "no": False}.get(v.lower()) for v in vals[1:])
            if winner is None or fx is None or fy is None:
                problems.append(f"{key}: verdict, fatal in X and fatal in Y must all be chosen (got {vals})")
            else:
                grades[key] = (winner, fx, fy)
    return grades, problems


# --- 6 and 7. results -------------------------------------------------------------------------------------------------


def _agrees_seat(a: SeatAnswer, p: Grade, judge: str | None = None) -> bool:
    """The Principal agrees with the judges when the fatal flag is the same and the total is within 1 point (or the
    usable verdict is the same); `judge` compares one judge alone."""
    if judge is not None:
        g = a.judges.get(judge)
        return g is not None and g.fatal == p.fatal and (abs(g.total - p.total) <= 1 or g.usable == p.usable)
    mean = a.judge_mean
    return (mean is not None and a.judge_fatal == p.fatal
            and (abs(mean - p.total) <= 1 or a.judges_usable == p.usable))


def _usable_by(a: SeatAnswer, judge: str) -> int:
    """1 when `judge`'s own grade of the answer is usable (the neutral judge's view in the seat rule)."""
    g = a.judges.get(judge)
    return int(g is not None and g.usable)


def _seat_agreement(mine: Sequence[SeatAnswer], rows: Mapping[str, str]) -> float | None:
    """The Principal's agreement with the judges on one domain's calibration rows; None without a graded one."""
    cal = [a for a in mine if rows.get(a.label) == "calibration" and a.principal is not None]
    return (sum(_agrees_seat(a, a.principal) for a in cal if a.principal) / len(cal)) if cal else None


def _trial_agreement(pairs: Mapping[str, PairTask], rows: Mapping[str, str]) -> float | None:
    """The Principal's agreement with the judges on the trial's calibration tasks; None without a graded one."""
    cal = [p for p in pairs.values() if rows.get(p.label) == "calibration" and p.principal is not None]
    agree = [p for p in cal if p.principal is not None and p.principal[0] == p.judge_winner
             and (p.principal[1], p.principal[2]) == (p.fatal["x"], p.fatal["y"])]
    return len(agree) / len(cal) if cal else None


def pending_full_grading(run: Run, rows: Mapping[str, str], *, seat: Mapping[str, SeatAnswer] | None = None,
                         pairs: Mapping[str, PairTask] | None = None) -> list[str]:
    """Rule 6, checked before anything is attributed to an engine: each D-49 domain, or the D-46 trial, whose
    calibration agreement is below 80% and which the Principal has not yet graded in full. Each line names the workbook
    to return. Hand grading stays blind until it is complete (notes_SETS.md section 4 item 6)."""
    out: list[str] = []
    for domain in sorted({a.domain for a in (seat or {}).values()}):
        mine = [a for a in (seat or {}).values() if a.domain == domain]
        rate = _seat_agreement(mine, rows)
        if rate is not None and rate < AGREEMENT_MIN and not all(a.principal is not None for a in mine):
            out.append(f"domain {domain}: calibration agreement {100 * rate:.0f}% is below 80%: grade the domain in "
                       f"full (atlas-admin eval workbook {run.run_id} --domain {domain}) and return it with the first "
                       "workbook")
    if pairs is not None:
        rate = _trial_agreement(pairs, rows)
        if rate is not None and rate < AGREEMENT_MIN and not all(p.principal is not None for p in pairs.values()):
            out.append(f"the trial: calibration agreement {100 * rate:.0f}% is below 80%: grade the whole trial in "
                       f"full (atlas-admin eval workbook {run.run_id} --all) and return it with the first workbook")
    return out


def _seat_results(run: Run, seat: dict[str, SeatAnswer], rows: Mapping[str, str]) -> dict[str, Any]:
    """Per domain: the tallies, the leads and the seat rule. results() has refused a pending domain already."""
    domains: dict[str, Any] = {}
    engines = run.contestants
    items = run.evset.by_id
    for domain in sorted({a.domain for a in seat.values()}):
        mine = [a for a in seat.values() if a.domain == domain]
        cal = [a for a in mine if rows.get(a.label) == "calibration" and a.principal is not None]
        rate = _seat_agreement(mine, rows)
        judge_rates = {}
        for j in engines:
            mine_j = [a for a in cal if j in a.judges]
            if mine_j:
                judge_rates[j] = round(sum(_agrees_seat(a, a.principal, j) for a in mine_j if a.principal)
                                       / len(mine_j), 4)
        rep: dict[str, Any] = {"answers": len(mine), "calibration_agreement": rate, "judge_agreement": judge_rates,
                               "drop_as_judge": sorted(j for j, r in judge_rates.items() if r < AGREEMENT_MIN)}
        by = {(a.engine, a.task): a for a in mine}
        tasks = sorted({a.task for a in mine})
        tallies = {}
        for e in engines:
            own = [t for t in tasks if items[t]["drafted_by"] != e]
            tallies[e] = Tally(len(own), sum(by[(e, t)].usable for t in own), sum(by[(e, t)].fatal for t in own))
        leads = {}
        for x in engines:
            if x == SEAT_INCUMBENT:
                continue
            neutral = next(e for e in engines if e not in (x, SEAT_INCUMBENT))
            common = [t for t in tasks if items[t]["drafted_by"] not in (x, SEAT_INCUMBENT)]
            lead = sum(by[(x, t)].usable for t in common) - sum(by[(SEAT_INCUMBENT, t)].usable for t in common)
            nlead = (sum(_usable_by(by[(x, t)], neutral) for t in common)
                     - sum(_usable_by(by[(SEAT_INCUMBENT, t)], neutral) for t in common))
            leads[x] = (lead, nlead)
        rep.update(status="decided", tallies={e: asdict(t) | {"rate": round(t.rate, 4)} for e, t in tallies.items()},
                   leads={x: {"lead": v[0], "neutral_lead": v[1]} for x, v in leads.items()},
                   **seat_rule(tallies, leads))
        domains[domain] = rep
    return {"domains": domains}


def _trial_results(run: Run, pairs: dict[str, PairTask], rows: Mapping[str, str], floor_met: bool | None,
                   answers: Mapping[str, Mapping[str, Any]]) -> dict[str, Any]:
    """The trial's counts and the promotion rule. results() has refused a pending trial already."""
    out: dict[str, Any] = {"calibration_agreement": _trial_agreement(pairs, rows)}
    wins = sum(1 for p in pairs.values() if p.winner_arm() == TRIAL_KEY)
    losses = sum(1 for p in pairs.values() if p.winner_arm() == TRIAL_INCUMBENT)
    ties = len(pairs) - wins - losses
    objective: dict[str, tuple[float, float]] = {}
    for cat in D46_OBJECTIVE:
        scores = []
        for arm in (TRIAL_KEY, TRIAL_INCUMBENT):
            total = 0.0
            for p in pairs.values():
                if p.category == cat:
                    rec = answers[run.label_of(p.task, arm)]
                    total += check_objective(run.item(p.task)["objective"], blinded_answer(rec))[0]
            scores.append(total)
        objective[cat] = (scores[0], scores[1])
    out.update(status="decided", **trial_rule(
        wins=wins, losses=losses, ties=ties,
        fatal_trial=sum(p.arm_fatal(TRIAL_KEY) for p in pairs.values()),
        fatal_incumbent=sum(p.arm_fatal(TRIAL_INCUMBENT) for p in pairs.values()),
        objective=objective, floor_met=floor_met))
    return out


def results(run: Run, grade_files: Sequence[Path], db: Path, *, floor_met: bool | None = None,
            euid: int | None = None) -> dict[str, Any]:
    """Rules 6 and 7: merge the returned workbooks; refuse (nothing unblinded) while an escalation or calibration row
    lacks the Principal's grade, or while a domain (D-49) or the trial (D-46) whose calibration agreement is below 80%
    is not yet graded in full; then unblind, apply the pre-registered rule, write results.json and record an
    `eval-unblind` row (counts and decisions only)."""
    source = verify_preregistration(run, db, required=False, euid=euid)
    grades: dict[str, Any] = {}
    problems: list[str] = []
    for f in grade_files:
        g, p = read_workbook(f, run)
        grades.update(g)  # a later workbook (a domain graded in full) overrides an earlier grade of the same row
        problems += [f"{f.name}: {x}" for x in p]
    if problems:
        raise EvalError("the returned workbook has problems; nothing is unblinded:\n" + "\n".join(problems))
    rows = principal_rows(run)
    missing = sorted(k for k in rows if k not in grades)
    if missing:
        raise EvalError(f"{len(missing)} row(s) still need the Principal's grade; nothing is unblinded: "
                        + ", ".join(missing[:20]) + (" ..." if len(missing) > 20 else ""))
    # Rule 6 before anything is attributed to an engine: a truncated answer is visibly cut off in the full-grading
    # workbook, so no truncation by engine, no results.json and no ledger row until no domain (or trial) is pending.
    seat: dict[str, SeatAnswer] = {}
    pairs: dict[str, PairTask] = {}
    if run.kind == SEAT_SET:
        seat = aggregate_seat(run)
        for lb, g in grades.items():
            if lb in seat:
                seat[lb].principal = g
        pending = pending_full_grading(run, rows, seat=seat)
    else:
        pairs = aggregate_pairs(run)
        by_label = {p.label: p for p in pairs.values()}
        for lb, g in grades.items():
            if lb in by_label:
                by_label[lb].principal = g
        pending = pending_full_grading(run, rows, pairs=pairs)
    if pending:
        raise EvalError("calibration agreement is below 80%, so grading goes on by hand, blind; nothing is "
                        "unblinded:\n" + "\n".join(pending))
    answers = {r["label"]: r for r in read_jsonl(run.answers_path)}
    truncations = truncation_summary(run, list(answers.values()))
    out: dict[str, Any] = {"run_id": run.run_id, "set": run.kind, "version": run.manifest["version"],
                           "preregistration": source, "principal_rows": len(grades),
                           "truncations": truncations,
                           "truncated_answers": sorted(f"{r['task']}:{run.writer(r['label'])}"
                                                       for r in answers.values() if r.get("truncated"))}
    if run.kind == SEAT_SET:
        out.update(_seat_results(run, seat, rows))
        out["todo"] = ("the Principal may close seat-test-sets once wave 2 is scored (only the Principal closes it)")
        summary = {d: {k: r.get(k) for k in ("status", "seat", "decision", "verifier_pass_required")}
                   for d, r in out["domains"].items()}
    else:
        out.update(_trial_results(run, pairs, rows, floor_met, answers))
        out["todo"] = "the Principal closes trial-task-set (T-30) with this result"
        summary = {k: out.get(k) for k in ("status", "decision", "wins", "decided", "wins_needed", "promise")}
    write_private(run.dir / "results.json", json.dumps(out, indent=1, sort_keys=True, default=str).encode("utf-8"))
    out["ledger_task_id"] = ledger_record(db, "eval-unblind", {"run_id": run.run_id, "set_sha256": run.manifest[
        "set_sha256"], "key_sha256": run.manifest["key_sha256"], "summary": summary}, euid=euid)
    run.log_event({"event": "unblinded", "ledger_task_id": out["ledger_task_id"]})
    return out


# --- CLI (atlas-admin eval ...) ---------------------------------------------------------------------------------------


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="atlas-admin eval", description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)

    def root_arg(s: argparse.ArgumentParser) -> None:
        s.add_argument("--eval-root", type=Path, default=DEFAULT_EVAL_ROOT,
                       help=f"root-only directory of sets, runs and keys (default {DEFAULT_EVAL_ROOT})")

    v = sub.add_parser("validate", help="strict schema check and the plan report (no write)")
    v.add_argument("set", type=Path)
    v.set_defaults(fn=_cmd_validate)
    f = sub.add_parser("freeze", help="freeze a checked set: run, blinding key, ledger record of hashes and rule")
    f.add_argument("set", type=Path)
    f.add_argument("--run-id", default=None)
    f.add_argument("--db", type=Path, default=None, help="override ATLAS_DB_PATH")
    root_arg(f)
    f.set_defaults(fn=_cmd_freeze)
    r = sub.add_parser("run", help="answers then judges, one engine at a time, through the Arbiter (resumable)")
    r.add_argument("run_id")
    r.add_argument("--phase", choices=("answers", "judges", "all"), default="all")
    r.add_argument("--engine", action="append", default=None, help="answer with these contestants only (repeat)")
    r.add_argument("--orch-url", default=None, help="default: ORCH_URL of /etc/atlas/orchestrator.env")
    r.add_argument("--admin-token-file", default=None, help="default: ORCH_ADMIN_TOKEN_FILE of orchestrator.env")
    r.add_argument("--trial-url", default=None, help=f"loopback URL of the trial engine ({TRIAL_KEY}, S-51)")
    r.add_argument("--trial-window", type=int, default=None, help="the trial engine's per-slot window, in tokens")
    r.add_argument("--db", type=Path, default=None)
    root_arg(r)
    r.set_defaults(fn=_cmd_run)
    w = sub.add_parser("workbook", help="the blinded scoring workbook for the Principal")
    w.add_argument("run_id")
    w.add_argument("--out", type=Path, default=None, help="default: <run>/workbook.xlsx")
    w.add_argument("--all", action="store_true", help="every row (D-46: the whole trial graded by hand)")
    w.add_argument("--domain", action="append", default=[], help="D-49: every answer of this domain (repeat)")
    root_arg(w)
    w.set_defaults(fn=_cmd_workbook)
    s = sub.add_parser("results", help="read the grades back, then unblind and apply the pre-registered rule")
    s.add_argument("run_id")
    s.add_argument("--grades", type=Path, action="append", required=True, help="a returned workbook (repeat)")
    s.add_argument("--floor-met", choices=("yes", "no"), default=None, help="D-46: the D-47 floor in T-31")
    s.add_argument("--db", type=Path, default=None)
    root_arg(s)
    s.set_defaults(fn=_cmd_results)
    return p


OWNER_UID = 0  # the sets, keys and runs are root's (rule 2)


def _require_root() -> None:
    if os.geteuid() != OWNER_UID:
        raise EvalError("run as root: the evaluation sets, keys and runs are root-only (sudo atlas-admin eval ...)")


def _ledger_path(args: argparse.Namespace) -> Path:
    """ATLAS_DB_PATH (or --db). It must exist: root creating it here would leave a ledger the service cannot write
    (Phase 2 step 2 creates it as the service's own; admin.py _root_owner_conflict)."""
    db = Path(args.db) if args.db else Settings.from_env().db_path
    if not db.is_file():
        raise EvalError(f"no ledger at {db}; it is created by Phase 2 step 2 (atlas-admin init-db as atlas)")
    return db


def _cmd_validate(args: argparse.Namespace) -> int:
    refuse_repository_path(args.set, "set")
    evset = load_set(args.set)
    print(f"valid: {args.set} {evset.kind} {evset.version}, {len(evset.items)} item(s), sha256 {evset.sha256}")
    for line in plan_report(evset):
        print(f"plan: {line}")
    return 0


def _cmd_freeze(args: argparse.Namespace) -> int:
    _require_root()
    db = _ledger_path(args)
    run = freeze(args.set, eval_root=args.eval_root, db=db, owner_uid=OWNER_UID, run_id=args.run_id)
    m = run.manifest
    print(f"frozen: run {run.run_id} ({m['set']} {m['version']}, {m['items']} items) set sha256 {m['set_sha256']}, "
          f"key sha256 {m['key_sha256']}, rule sha256 {m['rule_sha256']}; ledger row {m['ledger_task_id']}")
    print(f"next: atlas-admin eval run {run.run_id}")
    return 0


def _cmd_run(args: argparse.Namespace) -> int:
    _require_root()
    settings = Settings.from_env()
    run = Run.open(args.eval_root, args.run_id, owner_uid=OWNER_UID)
    verify_for_run(run, _ledger_path(args))
    engines = load_engines(settings.config_dir, settings.llama_port_base)
    prompts = SeatPrompts(load_personas(settings.config_dir, engines), load_domain_cards(settings.config_dir))
    orch_env = parse_env_file(settings.etc_dir / "orchestrator.env")
    url = args.orch_url or orch_env.get("ORCH_URL") or f"http://127.0.0.1:{orch_env.get('ORCH_PORT', '8800')}"
    token = read_admin_token(args.admin_token_file or orch_env.get("ORCH_ADMIN_TOKEN_FILE"))
    external: dict[str, tuple[str, int]] = {}
    wants_trial = run.kind == TRIAL_SET and args.phase != "judges" and (args.engine is None or TRIAL_KEY in args.engine)
    if wants_trial:
        if not args.trial_url or not args.trial_window:
            raise EvalError(f"the {TRIAL_KEY} arm needs --trial-url and --trial-window (the S-51 trial build's "
                            f"endpoint; UNVERIFIED until it exists); run the 0731 arm alone with --engine "
                            f"{TRIAL_INCUMBENT}")
        external[TRIAL_KEY] = (args.trial_url, args.trial_window)
    gate = OrchestratorGate(url, admin_token=token)
    client_for = client_factory(engines, env_dir=settings.engines_env_dir, external=external)
    if args.phase in ("answers", "all"):
        summary = run_answers(run, gate, client_for, prompts, engines=args.engine, external=set(external))
        print(json.dumps({"truncations": summary}, indent=1, sort_keys=True))
    if args.phase in ("judges", "all"):
        run_judges(run, gate, client_for)
        print(f"judged: next, atlas-admin eval workbook {run.run_id}")
    return 0


def _cmd_workbook(args: argparse.Namespace) -> int:
    _require_root()
    run = Run.open(args.eval_root, args.run_id, owner_uid=OWNER_UID)
    out = args.out or (run.dir / "workbook.xlsx")
    path, n = write_workbook(run, out, owner_uid=OWNER_UID, all_rows=args.all, domains=args.domain)
    print(f"workbook: {path} ({n} row(s), mode 600). Give a copy to the Principal; return it with atlas-admin eval "
          f"results {run.run_id} --grades <file>")
    return 0


def _cmd_results(args: argparse.Namespace) -> int:
    _require_root()
    run = Run.open(args.eval_root, args.run_id, owner_uid=OWNER_UID)
    floor = {"yes": True, "no": False}.get(args.floor_met or "")
    out = results(run, args.grades, _ledger_path(args), floor_met=floor)
    print(json.dumps(out, indent=1, sort_keys=True, default=str))
    return 0


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        return int(args.fn(args))
    except (EvalError, ConfigError, EngineError, ArbiterError, OSError, sqlite3.Error, KeyError, ValueError) as exc:
        print(f"atlas-admin eval: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
